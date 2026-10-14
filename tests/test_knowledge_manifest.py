"""Synthetic catalogue contract tests; no real regulatory dataset is substituted."""

import csv
import io
import zipfile
from dataclasses import replace
from functools import partial

import pytest
from openpyxl import Workbook

from crossborder_compliance.application.knowledge_manifest import (
    DIMENSIONS,
    FIELDS,
    ManifestError,
    ManifestPolicy,
    normalize_relative_path,
)
from crossborder_compliance.application.knowledge_manifest import (
    preview_manifest as build_preview,
)
from crossborder_compliance.infrastructure.knowledge_manifest_parser import NativeManifestParser

preview_manifest = partial(build_preview, parser=NativeManifestParser())

HEADERS = (
    "序号",
    "资料名称",
    "原文语言",
    "资料类型",
    "国家/地区",
    "研究专题",
    "发布机构",
    "文件格式",
    "相对路径",
)
ROW = (
    1,
    "Synthetic source",
    "lang-label",
    "category-label",
    "jurisdiction-label",
    "topic-label",
    "issuer-label",
    "PDF",
    "目录/source.pdf",
)


def policy(**kwargs):
    refs = {field: {ROW[FIELDS.index(field)]: (f"approved-{field}",)} for field in DIMENSIONS}
    return ManifestPolicy(dict(zip(FIELDS, HEADERS, strict=True)), refs, **kwargs)


def document(rows, kind="CSV", headers=HEADERS):
    if kind == "CSV":
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)
        return stream.getvalue().encode("utf-8-sig")
    book = Workbook()
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


@pytest.mark.parametrize("kind", ["CSV", "XLSX"])
def test_genuine_catalogue_preserves_all_columns_and_requires_file_inventory(kind):
    data = document([ROW], kind)
    preview = preview_manifest(data, kind=kind, policy=policy())
    assert len(preview.records) == 1 and len(dict(preview.records[0].values)) == 9
    assert preview.records[0].status == "MANIFEST_IMPORTED"
    assert dict(preview.records[0].metadata_refs)["jurisdiction"] == ("approved-jurisdiction",)
    assert (
        preview_manifest(data, kind=kind, policy=policy(), source_paths=()).records[0].status
        == "FILE_MISSING"
    )


@pytest.mark.parametrize("fmt", ["DOC", "DOCX", "HTML", "PDF", "TXT", "XHTML"])
def test_six_declared_formats_are_proposals_not_uploaded_documents(fmt):
    row = (*ROW[:7], fmt, f"folder/file.{fmt.lower()}")
    result = preview_manifest(document([row]), kind="CSV", policy=policy())
    assert not result.records[0].issues and result.records[0].status == "MANIFEST_IMPORTED"


@pytest.mark.parametrize(
    "path",
    [
        "../a.pdf",
        "/a.pdf",
        "C:\\a.pdf",
        "\\\\host\\a.pdf",
        "a//b.pdf",
        "a/./b.pdf",
        "a\x00.pdf",
        "https://a/b.pdf",
    ],
)
def test_unsafe_relative_paths_fail_closed(path):
    with pytest.raises(ManifestError, match="INVALID_RELATIVE_PATH"):
        normalize_relative_path(path)


def test_unicode_windows_full_path_matching_retains_duplicate_basenames():
    a = (*ROW[:8], "甲\\source.pdf")
    b = (2, *ROW[1:8], "乙/source.pdf")
    result = preview_manifest(
        document([a, b]),
        kind="CSV",
        policy=policy(),
        source_paths=("甲/source.pdf", "乙/source.pdf", "other.pdf"),
    )
    assert [r.status for r in result.records] == ["FILE_MATCHED", "FILE_MATCHED"]
    assert result.unlisted_files == ("other.pdf",)


def test_duplicate_paths_are_retained_and_all_flagged():
    result = preview_manifest(document([ROW, ROW]), kind="CSV", policy=policy())
    assert len(result.records) == 2
    assert all(any(i.code == "DUPLICATE_CANDIDATE" for i in r.issues) for r in result.records)


def test_case_policy_and_unicode_normalization_are_explicit():
    assert normalize_relative_path("e\u0301/x.PDF") == "é/x.PDF"
    assert normalize_relative_path("a/X.PDF", case_sensitive=False) == "a/x.pdf"
    result = preview_manifest(
        document([ROW]), kind="CSV", policy=policy(), source_paths=("目录/SOURCE.pdf",)
    )
    assert result.records[0].status == "FILE_MISSING"
    result = preview_manifest(
        document([ROW]),
        kind="CSV",
        policy=policy(case_sensitive_paths=False),
        source_paths=("目录/SOURCE.pdf",),
    )
    assert result.records[0].status == "FILE_MATCHED"


def test_ambiguous_archive_members_cannot_be_silently_selected():
    result = preview_manifest(
        document([ROW]), kind="CSV", policy=policy(), source_paths=(ROW[-1], ROW[-1])
    )
    assert any(i.code == "AMBIGUOUS_MATCH" for i in result.records[0].issues)


def test_unknown_metadata_is_unresolved_without_country_or_category_guessing():
    row = (*ROW[:4], "unknown", *ROW[5:])
    result = preview_manifest(document([row]), kind="CSV", policy=policy())
    assert any(
        i.field == "jurisdiction" and i.code == "METADATA_MAPPING_REQUIRED"
        for i in result.records[0].issues
    )


def test_format_mismatch_required_values_and_invalid_sequence_are_reported():
    row = (0, "", *ROW[2:7], "DOC", ROW[-1])
    issues = preview_manifest(document([row]), kind="CSV", policy=policy()).records[0].issues
    assert {i.code for i in issues} >= {
        "INVALID_SEQUENCE",
        "REQUIRED_FIELD_MISSING",
        "FORMAT_PATH_MISMATCH",
    }


def test_configurable_column_names_and_limits():
    renamed = tuple(f"field-{i}" for i in range(9))
    p = replace(policy(), columns=dict(zip(FIELDS, renamed, strict=True)))
    assert (
        preview_manifest(document([ROW], headers=renamed), kind="CSV", policy=p).records[0].status
        == "MANIFEST_IMPORTED"
    )
    with pytest.raises(ManifestError, match="MANIFEST_ROW_LIMIT"):
        preview_manifest(document([ROW, ROW]), kind="CSV", policy=policy(max_rows=1))
    with pytest.raises(ManifestError, match="MANIFEST_SIZE_EXCEEDED"):
        preview_manifest(document([ROW]), kind="CSV", policy=policy(max_bytes=1))


@pytest.mark.parametrize("kind", ["CSV", "XLSX"])
def test_missing_or_duplicated_column_is_rejected(kind):
    with pytest.raises(ManifestError, match="MISSING_OR_AMBIGUOUS_COLUMNS"):
        preview_manifest(
            document([ROW], kind, headers=(*HEADERS[:-1], HEADERS[0])), kind=kind, policy=policy()
        )


def test_xlsx_formula_is_never_evaluated_or_formalized():
    row = (1, '=HYPERLINK("https://invalid.example", "value")', *ROW[2:])
    result = preview_manifest(document([row], "XLSX"), kind="XLSX", policy=policy())
    assert any(i.code == "FORMULA_NOT_ALLOWED" for i in result.records[0].issues)


def test_workbook_expansion_and_external_entities_are_rejected():
    with pytest.raises(ManifestError, match="MANIFEST_EXPANSION_LIMIT"):
        preview_manifest(document([ROW], "XLSX"), kind="XLSX", policy=policy(max_expanded_bytes=10))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(
            "xl/workbook.xml", '<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x/>'
        )
    with pytest.raises(ManifestError, match="UNSAFE_WORKBOOK_XML"):
        preview_manifest(output.getvalue(), kind="XLSX", policy=policy())


def test_workbook_sheet_selection_is_required_when_ambiguous():
    stream = io.BytesIO()
    book = Workbook()
    book.active.append(HEADERS)
    book.active.append(ROW)
    book.create_sheet("other")
    book.save(stream)
    book.close()
    with pytest.raises(ManifestError, match="SHEET_SELECTION_REQUIRED"):
        preview_manifest(stream.getvalue(), kind="XLSX", policy=policy())
    assert (
        len(
            preview_manifest(
                stream.getvalue(), kind="XLSX", policy=policy(sheet_name="Sheet")
            ).records
        )
        == 1
    )


def test_invalid_configuration_and_corrupt_input_fail_closed():
    with pytest.raises(ManifestError, match="INVALID_COLUMN_MAPPING"):
        ManifestPolicy({}, {})
    with pytest.raises(ManifestError, match="INVALID_XLSX"):
        preview_manifest(b"not a workbook", kind="XLSX", policy=policy())
    with pytest.raises(ManifestError, match="INVALID_CSV"):
        preview_manifest(b"\xff\xff", kind="CSV", policy=policy())


@pytest.mark.parametrize("kind", ["CSV", "XLSX"])
def test_column_bound_is_enforced_before_building_record(kind):
    with pytest.raises(ManifestError, match="MANIFEST_COLUMN_LIMIT"):
        preview_manifest(document([ROW], kind), kind=kind, policy=policy(max_columns=8))


def test_utf16_entity_declaration_cannot_bypass_xml_guard():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(
            "xl/workbook.xml",
            '<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x/>'.encode("utf-16"),
        )
    with pytest.raises(ManifestError, match="UNSAFE_WORKBOOK_XML"):
        preview_manifest(output.getvalue(), kind="XLSX", policy=policy())


def test_multilingual_label_uses_explicit_backend_mapping():
    p = policy()
    refs = {**p.metadata, "language": {ROW[2]: ("language-a", "language-b")}}
    result = preview_manifest(document([ROW]), kind="CSV", policy=replace(p, metadata=refs))
    assert dict(result.records[0].metadata_refs)["language"] == ("language-a", "language-b")
