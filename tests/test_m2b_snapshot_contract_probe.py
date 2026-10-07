"""Empirical contract probe; PASS means the reported blocker was reproduced.

This is not a successful M2-B closure test. It intentionally asserts the
existing fail-closed response instead of weakening H's snapshot validation.
"""
# ruff: noqa: F401,F811 -- shared real PostgreSQL fixture graph
import io
import json
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from sqlalchemy import select
from test_m2b_document_inputs import host, upload
from test_phase1j_postgres import fixture, foundation_i, foundation_j

from crossborder_compliance.infrastructure.intake_composition import project_context
from crossborder_compliance.infrastructure.persistence import context_models as e
from crossborder_compliance.infrastructure.persistence import models as b
from crossborder_compliance.infrastructure.persistence.classification_repository import (
    PostgresFormalClassificationRepository,
)

pytestmark = pytest.mark.runtime_smoke


def document_bytes(value):
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Type"
    table.cell(1, 0).text = value
    table.cell(1, 1).text = "number"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_reused_data_item_cannot_satisfy_new_snapshot_inventory_pin(foundation_j, tmp_path):
    from test_m2a_intake import create

    f = foundation_j
    _, client, context, values = host(f, tmp_path)
    initial = create(client, values).json()
    project_id = initial["project_id"]
    base_url = f"/api/v1/projects/{project_id}/intake"
    media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    first_upload = upload(client, initial, document_bytes("first-value"), name="fields.docx", media=media)
    assert first_upload.status_code == 201, first_upload.text
    first_view = first_upload.json()
    first_document = first_view["items"][0]
    parsed = client.post(
        base_url + f"/documents/{first_document['document_version_id']}/parse",
        json={"expected_version": 2, "idempotency_key": str(uuid4())},
    )
    assert parsed.status_code == 200 and parsed.json()["items"][0]["parse_status"] == "COMPLETED", parsed.text
    first_confirm = client.post(base_url + "/confirm", json={"expected_version": 2})
    assert first_confirm.status_code == 200, first_confirm.text
    first_snapshot = first_confirm.json()["analysis_snapshot_id"]

    scoped_context = project_context(f["sf"], context, UUID(project_id))
    classifier = PostgresFormalClassificationRepository(f["sf"], scoped_context)
    with f["sf"]() as session:
        item = session.scalar(select(b.DataItemEntity).where(
            b.DataItemEntity.tenant_id == f["tenant"],
            b.DataItemEntity.project_id == project_id,
            b.DataItemEntity.name == "field",
        ))
        assert item is not None
        item_id = item.data_item_id
        original_traces = set(session.scalars(select(e.DataItemSourceTraceLinkEntity.source_trace_ref_id).where(
            e.DataItemSourceTraceLinkEntity.data_item_id == item_id,
        )))
    # H accepts the original snapshot; this is not a missing permission/config.
    classifier.prepare(UUID(project_id), UUID(first_snapshot), UUID(item_id), UUID(f["scheme_version"]))

    superseded = client.post(base_url + "/supersede", json={"expected_version": 2, "idempotency_key": str(uuid4())})
    assert superseded.status_code == 200, superseded.text
    replaced = client.post(base_url + "/documents", data={
        "expected_version": 3,
        "idempotency_key": str(uuid4()),
        "replace_document_id": first_document["document_id"],
    }, files={"file": ("fields.docx", document_bytes("later-value"), media)})
    assert replaced.status_code == 201, replaced.text
    second_document = replaced.json()["items"][0]
    assert second_document["document_id"] == first_document["document_id"]
    assert second_document["document_version_id"] != first_document["document_version_id"]
    second_parse = client.post(base_url + f"/documents/{second_document['document_version_id']}/parse", json={
        "expected_version": 4, "idempotency_key": str(uuid4()),
    })
    assert second_parse.status_code == 200 and second_parse.json()["items"][0]["parse_status"] == "COMPLETED", second_parse.text
    second_confirm = client.post(base_url + "/confirm", json={"expected_version": 4})
    assert second_confirm.status_code == 200, second_confirm.text
    second_snapshot = second_confirm.json()["analysis_snapshot_id"]
    assert second_snapshot != first_snapshot

    with f["sf"]() as session:
        pins = session.scalars(select(e.AnalysisSnapshotContextPinEntity).where(
            e.AnalysisSnapshotContextPinEntity.analysis_snapshot_id.in_([first_snapshot, second_snapshot]),
        )).all()
        inventory = {pin.analysis_snapshot_id: pin.data_inventory_version for pin in pins}
        detail = session.get(e.DataItemResolutionDetailEntity, item_id)
        all_items = session.scalars(select(b.DataItemEntity).where(b.DataItemEntity.project_id == project_id)).all()
        current_traces = set(session.scalars(select(e.DataItemSourceTraceLinkEntity.source_trace_ref_id).where(
            e.DataItemSourceTraceLinkEntity.data_item_id == item_id,
        )))
        assert inventory[first_snapshot] == detail.version == 1
        assert inventory[second_snapshot] == 2
        assert len(all_items) == 2  # Field + Type headers, no duplicated authoritative item.
        assert original_traces < current_traces

    with pytest.raises(LookupError, match="classification resource not found"):
        classifier.prepare(UUID(project_id), UUID(second_snapshot), UUID(item_id), UUID(f["scheme_version"]))
    # Old snapshot validation still succeeds, but the shared provenance set grew.
    classifier.prepare(UUID(project_id), UUID(first_snapshot), UUID(item_id), UUID(f["scheme_version"]))
    evidence = {
        "status": "BLOCKER_REPRODUCED_NOT_CLOSURE_PASS",
        "database": "real PostgreSQL",
        "project_id": project_id,
        "data_item_id": item_id,
        "first_snapshot_id": first_snapshot,
        "second_snapshot_id": second_snapshot,
        "first_document_version_id": first_document["document_version_id"],
        "second_document_version_id": second_document["document_version_id"],
        "first_inventory_version": inventory[first_snapshot],
        "second_inventory_version": inventory[second_snapshot],
        "retained_detail_version": detail.version,
        "first_snapshot_prepare": "ACCEPTED",
        "second_snapshot_prepare": "LookupError: classification resource not found",
        "original_trace_ids": sorted(original_traces),
        "current_trace_ids": sorted(current_traces),
        "provenance_links_versioned": False,
        "formal_data_items_duplicated": False,
    }
    target = Path("artifacts/m2b/snapshot_contract_probe.json")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(evidence, indent=2) + "\n")
