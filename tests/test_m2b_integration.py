"""Real binary + PostgreSQL intake, pinning, authorization and canonical workflow."""
# ruff: noqa: F401,F811
import io
from uuid import UUID, uuid4
import pytest
from sqlalchemy import select
from test_m2b_document_inputs import host, upload
from test_m2a_intake import create
from test_m2a_structured_intake import binding, confirm, start
from test_phase1j_postgres import fixture, foundation_i, foundation_j
from crossborder_compliance.domain.security import RepositoryContext
from crossborder_compliance.interfaces.api.dependencies import get_repository_context
from crossborder_compliance.infrastructure.persistence import models as b, document_models as d, decision_models as j
from crossborder_compliance.infrastructure.persistence.document_repositories import PostgresDocumentIntelligenceRepository

pytestmark = pytest.mark.runtime_smoke


def prepared(f, tmp_path, binary=b'Purpose: Generic\n', name='input.txt', media='text/plain'):
    binding(f)
    binding(f,code="GENERIC_FIELD",field="business_purpose",value_type="string")
    app,c,ctx,values=host(f,tmp_path)
    one=create(c,{**values,'data_volume':'3'}).json()
    uploaded=upload(c,one,binary,name=name,media=media)
    assert uploaded.status_code==201,uploaded.text
    doc=uploaded.json()['items'][0]
    url=f"/api/v1/projects/{one['project_id']}/intake"
    parsed=c.post(url+f"/documents/{doc['document_version_id']}/parse",json=dict(expected_version=2,idempotency_key=str(uuid4())))
    assert parsed.status_code==200 and parsed.json()['items'][0]['parse_status']=='COMPLETED',parsed.text
    record=c.get(url).json()
    return app,c,ctx,record,doc,url


@pytest.mark.parametrize("foundation_i", [{"no_data": True}], indirect=True)
def test_real_upload_confirm_snapshot_canonical_start_read_and_history(foundation_j,tmp_path):
    f=foundation_j; app,c,ctx,record,doc,url=prepared(f,tmp_path)
    confirmed=confirm(c,record); view=start(c,confirmed)
    assert view['status']=='COMPLETED',view
    assert view['project_id']==record['project_id'] and view['analysis_snapshot_id']==confirmed['analysis_snapshot_id']
    assert view['result_refs']['applicability']
    with f['sf']() as s:
        pin=s.scalar(select(d.AnalysisSnapshotParseRunPinEntity).where(d.AnalysisSnapshotParseRunPinEntity.analysis_snapshot_id==confirmed['analysis_snapshot_id']))
        assert pin.document_version_id==doc['document_version_id']
        snapshot=s.get(b.AnalysisSnapshotEntity,confirmed['analysis_snapshot_id']); before=dict(snapshot.provenance_json)
        old_run=pin.parse_run_id
    supersede=c.post(url+'/supersede',json=dict(expected_version=2,idempotency_key=str(uuid4())))
    assert supersede.status_code==200,supersede.text
    draft=c.get(url).json(); assert draft['version']==3
    later=upload(c,draft,b'Later: file\n',name='later.txt'); assert later.status_code==201
    assert start(c,confirmed)==view and c.get(f"/api/v1/workflows/{view['workflow_run_id']}").json()==view
    with f['sf']() as s:
        assert s.get(b.AnalysisSnapshotEntity,confirmed['analysis_snapshot_id']).provenance_json==before
        pin=s.scalar(select(d.AnalysisSnapshotParseRunPinEntity).where(d.AnalysisSnapshotParseRunPinEntity.analysis_snapshot_id==confirmed['analysis_snapshot_id']))
        assert pin.parse_run_id==old_run
    app.dependency_overrides[get_repository_context]=lambda:RepositoryContext.user(ctx.tenant_id,'other',ctx.permission.scopes)
    assert c.get(f"/api/v1/workflows/{view['workflow_run_id']}").status_code==404


def test_pending_document_prevents_confirmation_and_no_snapshot(foundation_j,tmp_path):
    f=foundation_j; _,c,_,values=host(f,tmp_path); binding(f)
    one=create(c,values).json(); attached=upload(c,one).json()
    result=c.post(f"/api/v1/projects/{one['project_id']}/intake/confirm",json={'expected_version':attached['intake_version']})
    assert result.status_code==422 and 'INSUFFICIENT_INPUT' in result.text
    record=c.get(f"/api/v1/projects/{one['project_id']}/intake").json()
    assert record['status']=='DRAFT' and record['analysis_snapshot_id'] is None


def test_client_cannot_inject_document_authority(foundation_j,tmp_path):
    _,c,_,values=host(foundation_j,tmp_path); one=create(c,values).json()
    url=f"/api/v1/projects/{one['project_id']}/intake/documents"
    for field in ('tenant_id','storage_ref','policy_id','formal_result'):
        response=c.post(url,data=dict(expected_version=1,idempotency_key=str(uuid4()),**{field:'bad'}),files={'file':('test.txt',b'Test','text/plain')})
        assert response.status_code==422


def test_current_authorization_for_missing_storage_and_read_policy(foundation_j,tmp_path):
    app,c,ctx,record,doc,url=prepared(foundation_j,tmp_path)
    app.state.document_object_storage=None
    assert c.get(url+'/documents').status_code==200
    assert c.get(url+'/documents/policy').json()['status']=='CAPABILITY_NOT_CONFIGURED'
    assert upload(c,record,name='later.txt').status_code==503
    app.dependency_overrides[get_repository_context]=lambda:RepositoryContext.user(uuid4(),'author',ctx.permission.scopes)
    assert c.get(url+'/documents/policy').status_code==404


def test_multi_item_never_bypasses_data_analysis(foundation_j,tmp_path):
    from docx import Document
    doc=Document(); table=doc.add_table(rows=2,cols=2)
    for cell,value in zip(table.rows[0].cells,['Field','Type'],strict=True): cell.text=value
    for cell,value in zip(table.rows[1].cells,['email','string'],strict=True): cell.text=value
    buffer=io.BytesIO(); doc.save(buffer)
    f=foundation_j
    _,c,_,record,_,_=prepared(f,tmp_path,buffer.getvalue(),'real.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    confirmed=confirm(c,record); view=start(c,confirmed)
    assert view['status']=='WARNING' and view['reason_codes']==['MULTI_SUBJECT_WORKFLOW_NOT_CONFIGURED'],view
    assert not view['result_refs'].get('applicability')
    with f['sf']() as s:
        snapshot=s.get(b.AnalysisSnapshotEntity,confirmed['analysis_snapshot_id'])
        assert snapshot.provenance_json['formal_workflow_plan']['mode']=='DATA_AWARE'
        for model in j.MODELS.values():
            assert s.scalar(select(model.result_id).where(model.analysis_snapshot_id==confirmed['analysis_snapshot_id'])) is None


def test_single_item_executes_existing_data_aware_services(foundation_j,tmp_path):
    from docx import Document
    doc=Document(); table=doc.add_table(rows=2,cols=1); table.cell(0,0).text='Field'; table.cell(1,0).text='email'
    buffer=io.BytesIO(); doc.save(buffer)
    f=foundation_j
    _,c,_,record,_,_=prepared(f,tmp_path,buffer.getvalue(),'one-column.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    confirmed=confirm(c,record); view=start(c,confirmed)
    assert view['status']=='COMPLETED',view
    assert view['result_refs']['classification'] and view['result_refs']['applicability'],view
    with f['sf']() as s:
        snapshot=s.get(b.AnalysisSnapshotEntity,confirmed['analysis_snapshot_id'])
        assert snapshot.provenance_json['formal_workflow_plan']['mode']=='DATA_AWARE'


@pytest.mark.parametrize('conflict',[False,True])
def test_actual_uploaded_manual_and_document_fact_merge_or_durable_review(foundation_j,tmp_path,conflict):
    from test_m2a_structured_intake import facts
    from crossborder_compliance.infrastructure.persistence import context_models as e
    f=foundation_j; binding(f,code='GENERIC_FIELD',field='business_purpose',value_type='string')
    _,c,_,values=host(f,tmp_path)
    one=create(c,{**values,'business_purpose':'Manual purpose'}).json()
    document=upload(c,one,b'Purpose: '+(b'Other purpose' if conflict else b'Manual purpose')+b'\n').json()['items'][0]
    url=f"/api/v1/projects/{one['project_id']}/intake"
    assert c.post(url+f"/documents/{document['document_version_id']}/parse",json=dict(expected_version=2,idempotency_key=str(uuid4()))).status_code==200
    confirmed=confirm(c,c.get(url).json()); rows=facts(f,one['project_id'])
    assert len(rows)==(2 if conflict else 1)
    if conflict:
        view=start(c,confirmed); assert view['status']=='REVIEW_REQUIRED' and view['review_id']
        with f['sf']() as s:
            assert s.get(b.ReviewTaskEntity,view['review_id']).status=='PENDING'
            assert s.scalar(select(e.ContextConflictEntity).where(e.ContextConflictEntity.project_id==one['project_id'],e.ContextConflictEntity.conflict_type=='BUSINESS_FACT_CONFLICT')) is not None
    else:
        assert rows[0].structured_provenance_json and rows[0].source_document_ids_json


def test_process_crash_after_acceptance_then_restart_has_one_parse(foundation_j,tmp_path):
    import json,os,subprocess,sys
    from pathlib import Path
    from sqlalchemy import func
    f=foundation_j; _,c,ctx,values=host(f,tmp_path); one=create(c,values).json()
    uploaded=upload(c,one).json(); document=uploaded['items'][0]
    spec=tmp_path/'trusted-worker.json'; spec.write_text(json.dumps(dict(tenant=f['tenant'],actor='author',scopes=sorted(ctx.permission.scopes),project=one['project_id'],document_version=document['document_version_id'],version=2,key=str(uuid4()))))
    root=Path(__file__).resolve().parents[1]
    child=subprocess.run([sys.executable,str(root/'tests/m2b_parse_worker.py'),str(spec)],cwd=root,env={**os.environ,'PYTHONPATH':str(root/'src')+os.pathsep+str(root)},capture_output=True,text=True)
    assert child.returncode==75,child.stderr
    with f['sf']() as s:
        task=s.scalar(select(d.DocumentParseTaskEntity).where(d.DocumentParseTaskEntity.document_version_id==document['document_version_id']))
        assert task.status=='ACCEPTED'
        task_id=task.task_id
    url=f"/api/v1/projects/{one['project_id']}/intake/documents/{document['document_version_id']}/parse"
    response=c.post(url,json=dict(expected_version=2,idempotency_key=str(uuid4())))
    assert response.status_code==200 and response.json()['items'][0]['parse_status']=='COMPLETED',response.text
    assert c.post(url,json=dict(expected_version=2,idempotency_key=str(uuid4()))).json()==response.json()
    with f['sf']() as s:
        assert s.get(d.DocumentParseTaskEntity,task_id).status=='COMPLETED'
        assert s.scalar(select(func.count()).select_from(b.DocumentParseRunEntity).where(b.DocumentParseRunEntity.document_version_id==document['document_version_id']))==1


def test_ungoverned_document_candidate_routes_review_before_formal_decisions(foundation_j,tmp_path):
    from crossborder_compliance.infrastructure.persistence import context_models as e
    f=foundation_j; binding(f); _,c,_,values=host(f,tmp_path)
    one=create(c,{**values,'data_volume':'3'}).json(); document=upload(c,one).json()['items'][0]
    url=f"/api/v1/projects/{one['project_id']}/intake"
    assert c.post(url+f"/documents/{document['document_version_id']}/parse",json=dict(expected_version=2,idempotency_key=str(uuid4()))).status_code==200
    confirmed=confirm(c,c.get(url).json()); view=start(c,confirmed)
    assert view['status']=='REVIEW_REQUIRED' and view['review_id'],view
    assert not view['result_refs'].get('applicability')
    with f['sf']() as s:
        assert s.scalar(select(e.BusinessFactEntity).where(e.BusinessFactEntity.project_id==one['project_id'],e.BusinessFactEntity.fact_type=='GENERIC_FIELD')) is None
        assert s.get(b.ReviewTaskEntity,view['review_id']).status=='PENDING'


def test_same_binary_wrong_project_is_explicit_conflict_without_identity_leak(foundation_j,tmp_path):
    f=foundation_j; _,c,_,values=host(f,tmp_path)
    a=create(c,values).json(); b_project=create(c,values).json()
    assert upload(c,a).status_code==201
    response=upload(c,b_project)
    assert response.status_code==409 and a['project_id'] not in response.text,response.text


def test_existing_conflict_takes_priority_over_multi_subject_capability_gap(foundation_j,tmp_path):
    from docx import Document
    doc=Document(); doc.add_paragraph('Purpose: Conflicting documented purpose')
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='Field';table.cell(1,0).text='email';table.cell(0,1).text='Type';table.cell(1,1).text='string'
    binary=io.BytesIO();doc.save(binary)
    _,c,_,record,_,_=prepared(foundation_j,tmp_path,binary.getvalue(),'conflicting-real.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    confirmed=confirm(c,record); view=start(c,confirmed)
    assert view['status']=='REVIEW_REQUIRED' and view['review_id'],view
    assert not view['result_refs'].get('applicability')
