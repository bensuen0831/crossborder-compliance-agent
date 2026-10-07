import io,json,sys
from uuid import UUID,uuid4
from docx import Document
from test_phase1j_migrations import seed_scope
from test_phase1e_postgres import MemoryStorage,CleanScan,Queue
from crossborder_compliance.config import get_settings
from crossborder_compliance.domain.security import RepositoryContext
from crossborder_compliance.application.document_services import DocumentIngestionService,DocumentParseService
from crossborder_compliance.infrastructure.document_parsers import default_native_parsers
from crossborder_compliance.infrastructure.persistence.document_repositories import PostgresDocumentIntelligenceRepository
from crossborder_compliance.infrastructure.persistence.context_repositories import PostgresContextResolutionRepository
from crossborder_compliance.application.context_services import DataItemNormalizationService
from sqlalchemy import text
engine,sf,ids=seed_scope(get_settings().database_url); tenant,project,*_=ids
ctx=RepositoryContext.user(UUID(tenant),'legacy-seed')
storage=MemoryStorage(); repo=PostgresDocumentIntelligenceRepository(sf,ctx)
doc=Document(); table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Field';table.cell(1,0).text='email';table.cell(0,1).text='Type';table.cell(1,1).text='string'
buffer=io.BytesIO();doc.save(buffer)
result=DocumentIngestionService(repo,storage,CleanScan()).ingest(tenant_id=UUID(tenant),project_id=UUID(project),filename='legacy-real.docx',mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',content=buffer.getvalue())
parser=DocumentParseService(repo,storage,default_native_parsers(),Queue());task=parser.request_parse(document_version_id=UUID(result['document_version_id']),idempotency_key=str(uuid4()));assert parser.process_task(UUID(task['task_id']))['status']=='COMPLETED'
context=PostgresContextResolutionRepository(sf,ctx)
run=context.start_context_run(UUID(project)); assert run['version']==1
DataItemNormalizationService(context).normalize(UUID(project),version=1)
with engine.begin() as c:
 details=c.execute(text('SELECT data_item_id,version,display_name FROM data_item_resolution_details ORDER BY data_item_id')).all()
 assert len(details)==2
 unknown_item=details[0][0]
 c.execute(text("DELETE FROM candidate_resolutions WHERE candidate_id IN (SELECT candidate_data_item_id FROM data_item_candidate_links WHERE data_item_id=:item)"),{'item':unknown_item})
 print(json.dumps(dict(details=[list(x) for x in details],unknown_item=unknown_item)))
