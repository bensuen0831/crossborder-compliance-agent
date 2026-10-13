"""Additive real HTTP/PG/DOCX consent + explicit typed cross-border flow proof."""
import io
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import select
from test_stage1_candidate_governance_http import (
    fixture, foundation_i, foundation_j, governed_http, prepare, decide,
)
from crossborder_compliance.infrastructure.persistence import models as b
from crossborder_compliance.infrastructure.persistence import context_models as e

pytestmark = pytest.mark.runtime_smoke


def test_real_http_typed_crossborder_flow_uses_confirmed_document_inventory(governed_http):
    f=governed_http
    destination=f['facts']['destination_locations'][0]
    flow=dict(input_key=str(uuid4()),data_item_canonical_name='Governed Data',
        source_jurisdiction_id=f['juri'],destination_jurisdiction_id=destination,
        direction='SOURCE_TO_DESTINATION',source_name='User source',destination_name='User destination')
    response=f['client'].put(f['path']+'/intake',json=dict(expected_version=f['version'],idempotency_key=uuid4().hex,
        facts={**f['facts'],'destination_locations':[destination],'data_flows':[flow]}))
    assert response.status_code==200,response.text
    f['version']=response.json()['version']
    document=Document();document.add_paragraph('Unstructured requirement PROVENANCE_MARKER_FLOW '+f['project'])
    table=document.add_table(rows=2,cols=1);table.cell(0,0).text='Governed Data';table.cell(1,0).text='Business measurement'
    binary=io.BytesIO();document.save(binary)
    upload=f['client'].post(f['path']+'/intake/documents',data=dict(expected_version=f['version'],idempotency_key=uuid4().hex),
        files={'file':('real-flow-input.docx',binary.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert upload.status_code==201,upload.text
    f['version']=upload.json()['intake_version']
    p=prepare(f);decision,_=decide(f,p)
    before=len(f['requests'])
    confirmed=f['client'].post(f['path']+'/intake/confirm',json=dict(expected_version=f['version'],candidate_decision_set_id=decision['candidate_decision_set_id']))
    assert confirmed.status_code==200,confirmed.text
    assert len(f['requests'])==before
    with f['sf']() as s:
        nodes=s.scalars(select(b.DataFlowNodeEntity).where(b.DataFlowNodeEntity.project_id==f['project'])).all()
        edges=s.scalars(select(b.DataFlowEdgeEntity).where(b.DataFlowEdgeEntity.project_id==f['project'])).all()
        assert len(nodes)==2 and len(edges)==1
        by_id={n.flow_node_id:n for n in nodes};edge=edges[0]
        assert by_id[edge.source_node_id].jurisdiction_id==f['juri']
        assert by_id[edge.target_node_id].jurisdiction_id==destination
        assert f['juri']!=destination
        links=s.scalars(select(b.DataItemFlowLinkEntity).where(b.DataItemFlowLinkEntity.flow_edge_id==edge.flow_edge_id)).all()
        assert len(links)==1
        item=s.get(b.DataItemEntity,links[0].data_item_id);assert item.project_id==f['project']
        detail=s.get(e.DataFlowEdgeDetailEntity,edge.flow_edge_id);assert detail.direction==flow['direction']
        snapshot=s.get(b.AnalysisSnapshotEntity,confirmed.json()['analysis_snapshot_id'])
        pin=s.scalar(select(e.AnalysisSnapshotContextPinEntity).where(e.AnalysisSnapshotContextPinEntity.analysis_snapshot_id==snapshot.analysis_snapshot_id))
        assert detail.version==pin.data_flow_version
        input_version=s.get(b.ProjectVersionEntity,snapshot.project_version_id)
        assert input_version.intake_json['data_flows'][0]['input_key']==flow['input_key']
        link_detail=s.get(e.DataItemFlowLinkDetailEntity,links[0].link_id)
        assert link_detail.source_trace_ids_json
