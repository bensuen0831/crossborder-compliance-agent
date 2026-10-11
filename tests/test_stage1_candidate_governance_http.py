"""Real PG/TCP/DOCX/HTTP-provider regression of the original R1 counterexample."""

import io
import json
from contextlib import ExitStack
from datetime import date
from types import SimpleNamespace
from uuid import UUID, uuid4

import pytest
from cryptography.fernet import Fernet
from docx import Document
from phase1kb_http_servers import local_provider
from sqlalchemy import func, select, text
from sqlalchemy.exc import IntegrityError
from test_m2a_intake import setup
from test_m2e_http import http_server
from test_m2e_vertical_postgres import server
from test_phase1c_model_registry import _ctx
from test_phase1k_a_configuration_postgres import publish as publish_metadata
from test_phase1kb_control_postgres import configured_model, provider_request, publish
from test_phase1l_b_postgres import fixture, foundation_i, foundation_j

from crossborder_compliance.application.model_control_services import ModelControlService
from crossborder_compliance.domain.llm_invocation import InvocationPurpose
from crossborder_compliance.domain.security import RepositoryContext
from crossborder_compliance.infrastructure.llm_gateway_http import check_endpoint
from crossborder_compliance.infrastructure.llm_provider_inspection import ProviderInspection
from crossborder_compliance.infrastructure.llm_secrets import EncryptedFileSecretStore
from crossborder_compliance.infrastructure.persistence import candidate_governance_models as g
from crossborder_compliance.infrastructure.persistence import context_models as e
from crossborder_compliance.infrastructure.persistence import document_models as d
from crossborder_compliance.infrastructure.persistence import models as b
from crossborder_compliance.infrastructure.persistence.config_admin_repositories import (
    PostgresGovernedArtifactAdminRepository,
)
from crossborder_compliance.infrastructure.persistence.metadata_repositories import (
    PostgresAdminMetadataRepository,
)
from crossborder_compliance.infrastructure.persistence.model_control import (
    PostgresModelControlRepository,
)
from crossborder_compliance.interfaces.api.dependencies import get_repository_context

pytestmark = pytest.mark.runtime_smoke


@pytest.fixture
def governed_http(tmp_path):
    f = foundation_i.__wrapped__(fixture.__wrapped__(), SimpleNamespace(param={"no_data": True}))
    f = foundation_j.__wrapped__(f, SimpleNamespace(param={}))
    app, facts = server(f, tmp_path)
    _, _, base_context, _ = setup(f)
    tenant = UUID(f["tenant"])
    admin = _ctx(tenant)
    metadata = PostgresAdminMetadataRepository(f["sf"], admin)
    note = metadata.create_definition(
        kind="BUSINESS_FACT_TYPE",
        code="DOCUMENT_NOTE_" + uuid4().hex,
        display_name="Governed document note",
    )
    publish_metadata(
        metadata,
        metadata.create_version(definition_id=UUID(note["definition_id"]), payload={}),
        True,
    )
    prompts = PostgresGovernedArtifactAdminRepository(f["sf"], admin, "prompts")
    prompt = publish_metadata(
        prompts,
        prompts.create_draft(
            code=uuid4().hex,
            display_name="Source-addressed candidates",
            payload={
                "template_text": "Return source-addressed candidates, never legal decisions.",
                "capability_requirement": ["STRUCTURED_OUTPUT"],
            },
        ),
    )
    prompts.bind_llm_invocation(
        UUID(prompt["version_id"]), InvocationPurpose.DOCUMENT_CANDIDATE_EXTRACTION
    )
    policy = metadata.create_definition(
        kind="LLM_INVOCATION_POLICY", code=uuid4().hex, display_name="Governed extraction"
    )
    publish_metadata(
        metadata,
        metadata.create_version(
            definition_id=UUID(policy["definition_id"]),
            payload={
                "allowed_triggers": {
                    "STANDARD": ["DOCUMENT_SEMANTIC_EXTRACTION_REQUIRED"],
                    "ENHANCED": ["DOCUMENT_SEMANTIC_EXTRACTION_REQUIRED"],
                },
                "max_models": 3,
                "candidate_fact_type_ids": [note["definition_id"]],
                "candidate_min_confidence": 0.95,
            },
        ),
        True,
    )
    requests = []

    def output(payload):
        text = payload["messages"][-1]["content"]
        if not text.startswith("{") or "source_nodes" not in json.loads(text):
            return {"ok": True}
        value = json.loads(text)
        node = next(n for n in value["source_nodes"] if "PROVENANCE_MARKER" in n["text"])
        requests.append(node["source_node_id"])
        return {
            "facts": [
                {
                    "source_node_id": node["source_node_id"],
                    "fact_type": note["code"],
                    "value": "ambiguous document note",
                    "confidence": 1,
                }
            ]
            if note["code"] in value["allowed_fact_types"]
            else [],
            "items": [
                {
                    "source_node_id": node["source_node_id"],
                    "name": "Generic Field",
                    "description": "document item",
                    "confidence": 1,
                }
            ],
            "nodes": [],
            "edges": [],
        }

    with ExitStack() as stack:
        provider_context = local_provider(("A1",), structured_response=output)
        url, calls = stack.enter_context(provider_context)
        models = PostgresModelControlRepository(f["sf"], admin)
        secrets = EncryptedFileSecretStore(
            tmp_path / "private-secrets", Fernet.generate_key(), tenant
        )
        service = ModelControlService(models, admin, secrets, check_endpoint)
        provider = publish(
            f["sf"],
            tenant,
            "PROVIDER",
            service.save_provider(
                provider_request(credential="ci-local-credential").model_copy(
                    update={"base_url": url}
                )
            ),
        )
        model = publish(
            f["sf"], tenant, "MODEL", configured_model(f["sf"], tenant, models, provider, "A1")
        )
        assert (
            ProviderInspection(models, secrets).model_test(UUID(model["deployment_id"]), "chat")[
                "status"
            ]
            == "HEALTHY"
        )
        user = RepositoryContext.user(
            tenant,
            "author",
            set(base_context.permission.scopes)
            | {
                "llm:invoke",
                "knowledge:retrieve",
                "document:read",
                "document:upload",
                "document:parse",
                "document:unlink",
                f"prompt:{prompt['definition_id']}:use",
            },
        )
        assert not user.permission.system
        app.dependency_overrides[get_repository_context] = lambda: user
        app.state.llm_secret_store_factory = lambda context: secrets
        c = stack.enter_context(http_server(app))
        c.timeout = 180
        response = c.post(
            "/api/v1/projects",
            json={
                "name": "R1 genuinely blank " + uuid4().hex,
                "idempotency_key": uuid4().hex,
                "facts": {"analysis_as_of_date": date.today().isoformat()},
            },
        )
        assert response.status_code == 201, response.text
        blank = response.json()
        project = blank["project_id"]
        assert blank["status"] == "DRAFT" and blank["analysis_snapshot_id"] is None
        for scope in ("TENANT", "PROJECT"):
            definition = metadata.create_definition(
                kind="MODEL_USAGE_POLICY", code=uuid4().hex, display_name=scope
            )
            publish_metadata(
                metadata,
                metadata.create_version(
                    definition_id=UUID(definition["definition_id"]),
                    payload={
                        "scope_type": scope,
                        "project_id": project if scope == "PROJECT" else None,
                        "default_mode": "INTERNAL_MODEL_ONLY",
                        "allowed_operations": ["structured_output"],
                        "allowed_model_ids": [model["model_id"]],
                        "allowed_provider_ids": [provider["provider_id"]],
                        "allowed_trust_levels": ["APPROVED"],
                        "allowed_data_boundaries": ["TENANT"],
                    },
                ),
                True,
            )
        path = f"/api/v1/projects/{project}"
        saved = c.put(
            path + "/intake",
            json={"expected_version": 1, "idempotency_key": uuid4().hex, "facts": facts},
        )
        assert saved.status_code == 200, saved.text
        doc = Document()
        doc.add_paragraph("Unstructured ambiguous note with PROVENANCE_MARKER " + project)
        binary = io.BytesIO()
        doc.save(binary)
        upload = c.post(
            path + "/intake/documents",
            data={"expected_version": saved.json()["version"], "idempotency_key": uuid4().hex},
            files={
                "file": (
                    "business-input.docx",
                    binary.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert upload.status_code == 201, upload.text
        attached = upload.json()
        version = attached["intake_version"]
        doc_id = attached["items"][0]["document_version_id"]
        native = c.post(
            path + f"/intake/documents/{doc_id}/parse",
            json={"expected_version": version, "idempotency_key": uuid4().hex},
        )
        assert native.status_code == 200, native.text
        assert native.json()["items"][0]["quality_status"] == "PASS" and not requests
        yield {
            **f,
            "client": c,
            "path": path,
            "project": project,
            "version": version,
            "fact_type": note["code"],
            "fact_definition": note,
            "requests": requests,
            "calls": calls,
            "stop_provider": lambda: provider_context.__exit__(None, None, None),
            "facts": facts,
            "doc_id": doc_id,
            "metadata": metadata,
            "model": model,
            "user": user,
            "app": app,
        }


def prepare(f, key=None):
    payload = {"expected_version": f["version"], "idempotency_key": key or uuid4().hex}
    response = f["client"].post(f["path"] + "/candidate-preparation", json=payload)
    assert response.status_code == 200, response.text
    value = response.json()
    assert value["status"] == "COMPLETED", value
    assert f["client"].get(f["path"] + "/candidate-preview").json() == value
    return value


def decide(f, p, action="ACCEPT", *, partial=False):
    decisions = []
    for item in p["candidates"]:
        chosen = action if item["candidate_type"] == "BUSINESS_FACT" else "ACCEPT"
        decision = {"preparation_candidate_id": item["preparation_candidate_id"], "action": chosen}
        if chosen == "EDIT":
            decision["edited_payload"] = {"normalized_value": "explicit corrected note"}
        decisions.append(decision)
    if partial:
        decisions = decisions[:1]
    payload = {
        "expected_version": f["version"],
        "candidate_preparation_id": p["candidate_preparation_id"],
        "candidate_universe_digest": p["candidate_universe_digest"],
        "decisions": decisions,
        "idempotency_key": uuid4().hex,
    }
    response = f["client"].post(f["path"] + "/candidate-decisions", json=payload)
    assert response.status_code == 200, response.text
    return response.json(), payload


@pytest.mark.parametrize("action", ["ACCEPT", "REJECT", "EDIT"])
def test_real_docx_candidate_decision_before_formalization_and_no_hidden_http(
    governed_http, action
):
    f = governed_http
    c = f["client"]
    p = prepare(f)
    assert len(f["requests"]) == 1
    assert any(
        i["candidate_origin"] == "LLM_STRUCTURED_OUTPUT" and i["confidence"] == 1
        for i in p["candidates"]
    )
    with f["sf"]() as s:
        assert (
            s.scalar(
                select(func.count())
                .select_from(e.BusinessFactEntity)
                .where(e.BusinessFactEntity.project_id == f["project"])
            )
            == 0
        )
        assert (
            s.scalar(
                select(func.count())
                .select_from(b.AnalysisSnapshotEntity)
                .join(
                    b.ProjectVersionEntity,
                    b.AnalysisSnapshotEntity.project_version_id
                    == b.ProjectVersionEntity.project_version_id,
                )
                .where(b.ProjectVersionEntity.project_id == f["project"])
            )
            == 0
        )
        audit = s.scalars(
            select(b.AuditEventEntity).where(
                b.AuditEventEntity.candidate_preparation_id == p["candidate_preparation_id"]
            )
        ).all()
        assert {a.event_type for a in audit} == {"LLM_ALLOWED", "LLM_COMPLETED"}
        assert all(
            a.invocation_scope == "DRAFT_PREPARATION"
            and a.workflow_run_id is None
            and a.analysis_snapshot_id is None
            for a in audit
        )
        assert all(a.provenance_json["actor_id"] == "author" for a in audit)
    blocked = c.post(f["path"] + "/intake/confirm", json={"expected_version": f["version"]})
    assert (
        blocked.status_code == 409 and blocked.json()["detail"] == "CANDIDATE_CONFIRMATION_REQUIRED"
    ), blocked.text
    decision, payload = decide(f, p, action)
    repeated = c.post(f["path"] + "/candidate-decisions", json=payload)
    assert repeated.json() == decision
    assert decision["submitted_by"] == "author"
    f["stop_provider"]()
    confirmed = c.post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    assert len(f["requests"]) == 1
    with f["sf"]() as s:
        formal = s.scalars(
            select(e.BusinessFactEntity).where(
                e.BusinessFactEntity.project_id == f["project"],
                e.BusinessFactEntity.fact_type == f["fact_type"],
            )
        ).all()
        if action == "REJECT":
            assert not formal
        else:
            assert len(formal) == 1 and formal[0].validation_status == "VALIDATED"
            assert formal[0].resolution_method == "USER_CONFIRMED_DOCUMENT_CANDIDATE"
            assert formal[0].normalized_value_json == (
                "explicit corrected note" if action == "EDIT" else "ambiguous document note"
            )
            mapping = s.scalar(
                select(e.CandidateResolutionEntity).where(
                    e.CandidateResolutionEntity.formal_object_id == formal[0].fact_id
                )
            )
            assert (
                mapping.resolved_by_type == "USER_CONFIRMATION" and mapping.resolved_at is not None
            )
            original = s.get(d.BusinessFactCandidateEntity, mapping.candidate_id)
            assert (
                original.normalized_value_json == "ambiguous document note"
                and original.validation_status == "UNVALIDATED"
            )
            assert mapping.source_trace_ids_json
        snapshot = s.get(b.AnalysisSnapshotEntity, confirmed.json()["analysis_snapshot_id"])
        assert snapshot.provenance_json["candidate_preparation_id"] == p["candidate_preparation_id"]
        assert (
            snapshot.provenance_json["candidate_decision_set_id"]
            == decision["candidate_decision_set_id"]
        )
        pins = s.scalars(
            select(g.CandidatePreparationDocumentPinEntity).where(
                g.CandidatePreparationDocumentPinEntity.candidate_preparation_id
                == p["candidate_preparation_id"]
            )
        ).all()
        snapshot_pins = s.scalars(
            select(d.AnalysisSnapshotParseRunPinEntity).where(
                d.AnalysisSnapshotParseRunPinEntity.analysis_snapshot_id
                == snapshot.analysis_snapshot_id
            )
        ).all()
        assert {(x.document_version_id, x.parse_run_id) for x in pins} == {
            (x.document_version_id, x.parse_run_id) for x in snapshot_pins
        }
        registry = s.scalars(
            select(g.CandidatePreparationRegistryPinEntity).where(
                g.CandidatePreparationRegistryPinEntity.candidate_preparation_id
                == p["candidate_preparation_id"]
            )
        ).all()
        from crossborder_compliance.infrastructure.persistence import metadata_models as m

        frozen = s.scalars(
            select(m.AnalysisSnapshotRegistryPinEntity).where(
                m.AnalysisSnapshotRegistryPinEntity.analysis_snapshot_id
                == snapshot.analysis_snapshot_id
            )
        ).all()
        tuples = lambda rows: {
            (r.pin_type, r.logical_key, r.object_id, r.version_id, r.version_no) for r in rows
        }
        assert tuples(registry) <= tuples(frozen)
        assert {
            "LLM_MODEL",
            "LLM_PROVIDER_VERSION",
            "LLM_USAGE_POLICY",
            "LLM_INVOCATION_POLICY",
            "LLM_PROMPT",
            "LLM_FACT_TYPE",
            "LLM_SELECTION",
        } <= {r.pin_type for r in registry}


def test_exact_universe_partial_unknown_and_idempotency(governed_http):
    f = governed_http
    c = f["client"]
    key = uuid4().hex
    p = prepare(f, key)
    assert prepare(f, key) == p and prepare(f) == p and len(f["requests"]) == 1
    bad = c.post(
        f["path"] + "/candidate-preparation",
        json={"expected_version": f["version"] + 1, "idempotency_key": key},
    )
    assert bad.status_code == 409
    decision, payload = decide(f, p, partial=True)
    blocked = c.post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert (
        blocked.status_code == 409 and blocked.json()["detail"] == "CANDIDATE_CONFIRMATION_REQUIRED"
    )
    unknown = {
        **payload,
        "idempotency_key": uuid4().hex,
        "decisions": [{"preparation_candidate_id": str(uuid4()), "action": "ACCEPT"}],
    }
    denied = c.post(f["path"] + "/candidate-decisions", json=unknown)
    assert denied.status_code == 409 and denied.json()["detail"] == "INVALID_CANDIDATE_DECISION"
    different = {**payload, "decisions": [dict(payload["decisions"][0], action="REJECT")]}
    assert c.post(f["path"] + "/candidate-decisions", json=different).status_code == 409
    spoof = {**payload, "idempotency_key": uuid4().hex, "actor_id": "another-user"}
    assert c.post(f["path"] + "/candidate-decisions", json=spoof).status_code == 422


@pytest.mark.parametrize("change", ["DRAFT", "DOCUMENT", "AI"])
def test_decision_is_stale_after_input_universe_changes(governed_http, change):
    f = governed_http
    c = f["client"]
    p = prepare(f)
    decision, _ = decide(f, p)
    if change == "DOCUMENT":
        doc = Document()
        doc.add_paragraph("Replacement material PROVENANCE_MARKER")
        binary = io.BytesIO()
        doc.save(binary)
        modified = c.post(
            f["path"] + "/intake/documents",
            data={"expected_version": f["version"], "idempotency_key": uuid4().hex},
            files={
                "file": (
                    "replacement.docx",
                    binary.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert modified.status_code == 201, modified.text
        version = modified.json()["intake_version"]
    else:
        facts = dict(f["facts"])
        if change == "AI":
            facts["ai_model_preference"] = {
                "usage_mode": "ENHANCED",
                "selection_mode": "SINGLE",
                "selected_model_ids": [f["model"]["model_id"]],
            }
        else:
            facts["business_purpose"] = "Explicitly changed future draft purpose"
        modified = c.put(
            f["path"] + "/intake",
            json={"expected_version": f["version"], "idempotency_key": uuid4().hex, "facts": facts},
        )
        assert modified.status_code == 200, modified.text
        version = modified.json()["version"]
    denied = c.post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": version,
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert denied.status_code == 409 and denied.json()["detail"] == "CANDIDATE_PREPARATION_STALE", (
        denied.text
    )
    assert (
        c.post(
            f["path"] + "/candidate-preparation",
            json={"expected_version": f["version"], "idempotency_key": uuid4().hex},
        ).status_code
        == 409
    )


def test_prepared_candidates_decisions_and_audit_are_append_only(governed_http):
    f = governed_http
    p = prepare(f)
    decision, _ = decide(f, p)
    mutations = [
        (
            "candidate_preparations",
            "candidate_preparation_id",
            p["candidate_preparation_id"],
            "created_by",
        ),
        (
            "candidate_decision_sets",
            "candidate_decision_set_id",
            decision["candidate_decision_set_id"],
            "submitted_by",
        ),
    ]
    with f["sf"]() as s:
        original = s.scalar(
            select(d.BusinessFactCandidateEntity)
            .join(
                g.CandidatePreparationCandidateEntity,
                g.CandidatePreparationCandidateEntity.candidate_id
                == d.BusinessFactCandidateEntity.fact_id,
            )
            .where(
                g.CandidatePreparationCandidateEntity.candidate_preparation_id
                == p["candidate_preparation_id"]
            )
        )
        assert original is not None
        mutations.append(
            ("business_fact_candidates", "fact_id", original.fact_id, "extraction_method")
        )
        row = s.scalar(
            select(g.CandidateDecisionEntity).where(
                g.CandidateDecisionEntity.candidate_decision_set_id
                == decision["candidate_decision_set_id"]
            )
        )
        mutations.append(
            ("candidate_decisions", "candidate_decision_id", row.candidate_decision_id, "action")
        )
        audit = s.scalar(
            select(b.AuditEventEntity).where(
                b.AuditEventEntity.candidate_preparation_id == p["candidate_preparation_id"]
            )
        )
        mutations.append(("audit_events", "audit_event_id", audit.audit_event_id, "event_type"))
    for table, key, identity, field in mutations:
        for statement in (
            f"UPDATE {table} SET {field}=:value WHERE {key}=:id",
            f"DELETE FROM {table} WHERE {key}=:id",
        ):
            with pytest.raises(IntegrityError, match="IMMUTABLE"):
                with f["sf"]() as s, s.begin():
                    s.execute(
                        text(statement),
                        {"id": identity, "value": "EDIT" if field == "action" else "spoof"},
                    )
    # A second human submission creates a new set; the first set is untouched.
    second, _ = decide(f, p, "REJECT")
    assert second["candidate_decision_set_id"] != decision["candidate_decision_set_id"]


def test_revoked_model_configuration_cannot_be_substituted(governed_http):
    from crossborder_compliance.infrastructure.persistence import metadata_models as m

    f = governed_http
    p = prepare(f)
    decision, _ = decide(f, p)
    with f["sf"]() as s, s.begin():
        row = s.get(m.ModelDeploymentEntity, f["model"]["deployment_id"])
        row.status = "DISABLED"
    denied = f["client"].post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert (
        denied.status_code == 409
        and denied.json()["detail"] == "CANDIDATE_PREPARATION_REVALIDATION_FAILED"
    ), denied.text
    assert len(f["requests"]) == 1


def test_external_channel_uses_same_preparation_decision_and_confirm_owner(governed_http):
    from test_m2e_http import authenticated

    from crossborder_compliance.domain.security import RepositoryContext

    f = governed_http
    c = f["client"]
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "sdk/python/src"))
    from crossborder_agent.client import AgentClient

    admin = RepositoryContext.user(
        UUID(f["tenant"]), "integration-administrator", {"integration:manage"}
    )
    f["app"].dependency_overrides[get_repository_context] = lambda: admin
    auth, identity = authenticated(c)
    client_id = identity["client"]["client_id"]
    binding = c.put(
        f"/api/v1/admin/integration-clients/{client_id}/project-binding",
        headers={"Idempotency-Key": uuid4().hex},
        json={
            "project_id": f["project"],
            "status": "ACTIVE",
            "expected_version": identity["client"]["record_version"],
        },
    )
    assert binding.status_code == 200, binding.text
    path = "/api/v1/external/projects/" + f["project"]
    sdk = AgentClient(str(c.base_url), client=c)
    sdk.use_service_credential(auth["Authorization"].removeprefix("Bearer "))
    p = sdk.prepare_candidates(f["project"], f["version"], uuid4().hex)
    assert sdk.candidate_preview(f["project"]) == p
    assert len(f["requests"]) == 1
    blocked = c.post(
        path + "/intake/confirm",
        headers={**auth, "Idempotency-Key": uuid4().hex},
        json={"expected_version": f["version"]},
    )
    assert (
        blocked.status_code == 409
        and blocked.json()["error_code"] == "CANDIDATE_CONFIRMATION_REQUIRED"
    )
    payload = {
        "expected_version": f["version"],
        "candidate_preparation_id": p["candidate_preparation_id"],
        "candidate_universe_digest": p["candidate_universe_digest"],
        "decisions": [
            {"preparation_candidate_id": item["preparation_candidate_id"], "action": "ACCEPT"}
            for item in p["candidates"]
        ],
    }
    resolved = sdk.resolve_candidates(f["project"], payload, uuid4().hex)
    assert resolved["submitted_by"] == "integration:" + client_id
    f["stop_provider"]()
    confirmed = sdk.confirm(
        f["project"],
        f["version"],
        uuid4().hex,
        candidate_decision_set_id=resolved["candidate_decision_set_id"],
    )
    assert confirmed["status"] == "CONFIRMED"
    other, _ = authenticated(c)
    assert c.get(path + "/candidate-preview", headers=other).status_code == 404


def test_reparse_without_version_change_invalidates_prepared_document_pins(governed_http):
    from crossborder_compliance.application.document_services import DocumentParseService
    from crossborder_compliance.infrastructure.document_parsers import default_native_parsers
    from crossborder_compliance.infrastructure.intake_composition import project_context
    from crossborder_compliance.infrastructure.persistence.document_repositories import (
        PostgresDocumentIntelligenceRepository,
    )

    f = governed_http
    p = prepare(f)
    decision, _ = decide(f, p)

    # Intake's completed native task correctly replays without a new run. A
    # genuine new canonical parser task exercises actual parse-universe drift.
    class Queue:
        def enqueue(self, **kwargs):
            pass

    context = project_context(f["sf"], f["user"], UUID(f["project"]))
    service = DocumentParseService(
        PostgresDocumentIntelligenceRepository(f["sf"], context),
        f["app"].state.document_object_storage,
        default_native_parsers(),
        Queue(),
    )
    task = service.request_parse(
        document_version_id=UUID(f["doc_id"]),
        idempotency_key=uuid4().hex,
        parser_profile_id="native-v1",
    )
    service.process_task(UUID(task["task_id"]))
    blocked = f["client"].post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert (
        blocked.status_code == 409 and blocked.json()["detail"] == "CANDIDATE_PREPARATION_STALE"
    ), blocked.text


def test_successor_requires_new_preparation_and_leaves_previous_snapshot_and_fact_unchanged(
    governed_http,
):
    from crossborder_compliance.infrastructure.persistence.candidate_governance import (
        content_digest,
    )

    f = governed_http
    c = f["client"]
    p = prepare(f)
    decision, _ = decide(f, p)
    first = c.post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert first.status_code == 200, first.text
    with f["sf"]() as s:
        old = s.get(b.AnalysisSnapshotEntity, first.json()["analysis_snapshot_id"])
        snapshot_digest = content_digest(
            {column.name: getattr(old, column.name) for column in old.__table__.columns}
        )
        fact = s.scalar(
            select(e.BusinessFactEntity).where(
                e.BusinessFactEntity.project_id == f["project"],
                e.BusinessFactEntity.fact_type == f["fact_type"],
            )
        )
        fact_id = fact.fact_id
        fact_digest = content_digest(
            {column.name: getattr(fact, column.name) for column in fact.__table__.columns}
        )
    successor = c.post(
        f["path"] + "/intake/supersede",
        json={"expected_version": f["version"], "idempotency_key": uuid4().hex},
    )
    assert successor.status_code == 200, successor.text
    f["version"] = successor.json()["intake_version"]
    stale = c.post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert stale.status_code == 409 and stale.json()["detail"] == "CANDIDATE_PREPARATION_STALE"
    prepared = prepare(f)
    edited, _ = decide(f, prepared, "EDIT")
    second = c.post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": edited["candidate_decision_set_id"],
        },
    )
    assert second.status_code == 200, second.text
    assert second.json()["analysis_snapshot_id"] != first.json()["analysis_snapshot_id"]
    with f["sf"]() as s:
        old = s.get(b.AnalysisSnapshotEntity, first.json()["analysis_snapshot_id"])
        assert (
            content_digest(
                {column.name: getattr(old, column.name) for column in old.__table__.columns}
            )
            == snapshot_digest
        )
        fact = s.get(e.BusinessFactEntity, fact_id)
        assert (
            content_digest(
                {column.name: getattr(fact, column.name) for column in fact.__table__.columns}
            )
            == fact_digest
        )
        new = s.get(b.AnalysisSnapshotEntity, second.json()["analysis_snapshot_id"])
        assert new.analysis_as_of_date == old.analysis_as_of_date
        assert (
            new.provenance_json["candidate_preparation_id"]
            != old.provenance_json["candidate_preparation_id"]
        )


def test_explicit_structured_input_is_not_reinferred_by_document_model(governed_http):
    f = governed_http
    row = f["metadata"].create_version(
        definition_id=UUID(f["fact_definition"]["definition_id"]),
        payload={
            "structured_intake_binding": {"field_path": "business_purpose", "value_type": "string"}
        },
    )
    publish_metadata(f["metadata"], row, True)
    preview = prepare(f)
    assert not any(
        item["candidate_type"] == "BUSINESS_FACT"
        and item["proposed_payload"].get("fact_type") == f["fact_type"]
        for item in preview["candidates"]
    )
    decision, _ = decide(f, preview)
    confirmed = f["client"].post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    with f["sf"]() as s:
        formal = s.scalar(
            select(e.BusinessFactEntity).where(
                e.BusinessFactEntity.project_id == f["project"],
                e.BusinessFactEntity.fact_type == f["fact_type"],
            )
        )
        assert formal.normalized_value_json == f["facts"]["business_purpose"]
        assert (
            formal.structured_provenance_json
            and formal.structured_provenance_json[0]["source_type"] == "USER_INPUT"
        )


def test_confirm_transfers_prepared_fact_type_version_after_future_publication(governed_http):
    from crossborder_compliance.infrastructure.persistence import metadata_models as m

    f = governed_http
    p = prepare(f)
    decision, _ = decide(f, p)
    with f["sf"]() as s:
        pinned = s.scalar(
            select(g.CandidatePreparationRegistryPinEntity).where(
                g.CandidatePreparationRegistryPinEntity.candidate_preparation_id
                == p["candidate_preparation_id"],
                g.CandidatePreparationRegistryPinEntity.pin_type == "LLM_FACT_TYPE",
            )
        )
        pinned_id = pinned.version_id
    future = publish_metadata(
        f["metadata"],
        f["metadata"].create_version(
            definition_id=UUID(f["fact_definition"]["definition_id"]),
            payload={"display_hint": "Future metadata presentation"},
        ),
        True,
    )
    assert future["version_id"] != pinned_id
    f["stop_provider"]()
    confirmed = f["client"].post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    with f["sf"]() as s:
        actual = s.scalar(
            select(m.AnalysisSnapshotRegistryPinEntity).where(
                m.AnalysisSnapshotRegistryPinEntity.analysis_snapshot_id
                == confirmed.json()["analysis_snapshot_id"],
                m.AnalysisSnapshotRegistryPinEntity.pin_type == "LLM_FACT_TYPE",
            )
        )
        assert actual.version_id == pinned_id


def test_candidate_added_after_review_fails_closed_without_partial_formalization(governed_http):
    f = governed_http
    p = prepare(f)
    decision, _ = decide(f, p)
    with f["sf"]() as s, s.begin():
        original = s.get(
            d.BusinessFactCandidateEntity,
            next(
                c["candidate_id"] for c in p["candidates"] if c["candidate_type"] == "BUSINESS_FACT"
            ),
        )
        fields = {
            column.name: getattr(original, column.name) for column in original.__table__.columns
        }
        fields["fact_id"] = str(uuid4())
        s.add(d.BusinessFactCandidateEntity(**fields))
    response = f["client"].post(
        f["path"] + "/intake/confirm",
        json={
            "expected_version": f["version"],
            "candidate_decision_set_id": decision["candidate_decision_set_id"],
        },
    )
    assert response.status_code == 409 and response.json()["detail"] == "CANDIDATE_UNIVERSE_CHANGED"
    with f["sf"]() as s:
        assert (
            s.scalar(
                select(func.count())
                .select_from(b.AnalysisSnapshotEntity)
                .join(
                    b.ProjectVersionEntity,
                    b.ProjectVersionEntity.project_version_id
                    == b.AnalysisSnapshotEntity.project_version_id,
                )
                .where(b.ProjectVersionEntity.project_id == f["project"])
            )
            == 0
        )


def test_database_refuses_decision_candidate_from_other_preparation(governed_http):
    f = governed_http
    old = prepare(f)
    modified = f["client"].put(
        f["path"] + "/intake",
        json={
            "expected_version": f["version"],
            "idempotency_key": uuid4().hex,
            "facts": {**f["facts"], "business_purpose": "New governed draft"},
        },
    )
    assert modified.status_code == 200
    f["version"] = modified.json()["version"]
    current = prepare(f)
    assert current["candidate_preparation_id"] != old["candidate_preparation_id"]
    with pytest.raises(IntegrityError):
        with f["sf"]() as s, s.begin():
            prep = s.get(g.CandidatePreparationEntity, current["candidate_preparation_id"])
            decision_set = g.CandidateDecisionSetEntity(
                candidate_decision_set_id=str(uuid4()),
                tenant_id=f["tenant"],
                project_id=f["project"],
                project_version_id=prep.project_version_id,
                candidate_preparation_id=prep.candidate_preparation_id,
                candidate_universe_digest=prep.candidate_universe_digest,
                decision_payload_digest="a" * 64,
                submitted_by="author",
                idempotency_key=uuid4().hex,
                request_digest="b" * 64,
            )
            s.add(decision_set)
            s.flush()
            s.add(
                g.CandidateDecisionEntity(
                    candidate_decision_id=str(uuid4()),
                    tenant_id=f["tenant"],
                    candidate_preparation_id=prep.candidate_preparation_id,
                    candidate_decision_set_id=decision_set.candidate_decision_set_id,
                    preparation_candidate_id=old["candidates"][0]["preparation_candidate_id"],
                    action="ACCEPT",
                    edited_payload_json=None,
                    decision_digest="c" * 64,
                )
            )
            s.flush()
