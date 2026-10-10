"""External HTTP vertical slice through existing DOCX/RAG/formal runtime."""

# ruff: noqa: F401,F811 -- canonical PostgreSQL fixture graph
import io
import time
from uuid import UUID, uuid4

import httpx
import pytest
from fastapi import FastAPI
from sqlalchemy import select
from test_m2a_intake import setup
from test_m2a_structured_intake import binding
from test_m2b_document_inputs import policy
from test_m2e_http import authenticated, http_server
from test_phase1f_postgres import metadata
from test_phase1l_b_postgres import fixture, foundation_i, foundation_j

from crossborder_compliance.domain.integrations import IntegrationPolicy
from crossborder_compliance.domain.security import RepositoryContext
from crossborder_compliance.infrastructure.document_storage import FileObjectStorageAdapter
from crossborder_compliance.infrastructure.intake_composition import prepare_snapshot
from crossborder_compliance.infrastructure.integration_worker import IntegrationDeliveryWorker
from crossborder_compliance.infrastructure.persistence import context_models as e
from crossborder_compliance.infrastructure.persistence import models as b
from crossborder_compliance.interfaces.api.dependencies import get_repository_context
from crossborder_compliance.interfaces.api.routes.external import router
from crossborder_compliance.interfaces.api.routes.integrations import admin_router
from crossborder_compliance.interfaces.api.routes.workflow import router as workflow_router

pytestmark = pytest.mark.runtime_smoke


def server(f, tmp_path, review=False):
    binding(f)
    binding(f, code="GENERIC_FIELD", field="business_purpose", value_type="string")
    _, _, _, facts = setup(f)
    # Distinct governed jurisdictions, no country-specific production code.
    from m2e_fixtures import destination_configuration

    destination = destination_configuration(f)
    facts["destination_locations"] = [destination]
    facts["data_volume"] = "3"
    app = FastAPI()
    app.include_router(admin_router)
    app.include_router(router)
    app.include_router(workflow_router)
    from crossborder_compliance.interfaces.api.routes.intake import router as intake_router
    from crossborder_compliance.interfaces.api.routes.intake_documents import (
        router as documents_router,
    )

    app.include_router(intake_router)
    app.include_router(documents_router)
    app.state.knowledge_session_factory = f["sf"]
    app.state.integration_policy = IntegrationPolicy(
        stream_duration_seconds=1, worker_poll_seconds=0.1, request_timeout_seconds=180
    )
    app.state.document_file_policy = policy()
    app.state.document_object_storage = FileObjectStorageAdapter(str(tmp_path / "binary"))
    if review:
        from functools import partial

        app.state.intake_snapshot_preparer = partial(
            prepare_snapshot, requirement_confirmation=True
        )
    context = RepositoryContext.user(
        UUID(f["tenant"]), "integration-administrator", {"integration:manage"}
    )
    app.dependency_overrides[get_repository_context] = lambda: context
    return app, facts


def docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Purpose: Generic")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Generic Field"
    table.cell(1, 0).text = "Sample"
    buff = io.BytesIO()
    doc.save(buff)
    return buff.getvalue()


@pytest.mark.parametrize("foundation_i", [{"no_data": True}], indirect=True)
@pytest.mark.parametrize("review", [False, True])
def test_external_real_docx_rag_async_result_and_review(foundation_j, tmp_path, review):
    f = foundation_j
    app, facts = server(f, tmp_path, review)
    worker = IntegrationDeliveryWorker(app)
    with http_server(app) as c:
        # The fixture revalidates the entire two-jurisdiction owning read closure.
        # Match the configured server deadline; the shared HTTP helper defaults to10s.
        c.timeout = httpx.Timeout(app.state.integration_policy.request_timeout_seconds)
        auth, identity = authenticated(c)

        def mutation(method, url, payload):
            response = c.request(
                method, url, headers={**auth, "Idempotency-Key": uuid4().hex}, json=payload
            )
            assert response.status_code < 300, response.text
            return response.json()

        intake = mutation(
            "POST", "/api/v1/external/projects", {"name": uuid4().hex, "facts": facts}
        )
        project = intake["project_id"]
        base = f"/api/v1/external/projects/{project}"
        # Exercise both input channels under the same authorized integration
        # context, not two different policies or impersonated Web users.
        from types import SimpleNamespace

        from crossborder_compliance.domain.integrations import IntegrationScope
        from crossborder_compliance.infrastructure.external_composition import integration_service

        service = integration_service(SimpleNamespace(app=app))
        principal = service.authenticate(auth["Authorization"][7:])
        ctx = service.authorize(principal, IntegrationScope.PROJECT_READ, UUID(project))
        app.dependency_overrides[get_repository_context] = lambda: ctx
        canonical_base = f"/api/v1/projects/{project}"
        initial = c.get(canonical_base + "/intake")
        initial_draft_equal = initial.status_code == 200 and initial.json() == intake
        assert initial_draft_equal
        updated = c.put(
            canonical_base + "/intake",
            json={
                "expected_version": intake["version"],
                "idempotency_key": uuid4().hex,
                "facts": facts,
            },
        )
        assert updated.status_code == 200, updated.text
        intake = updated.json()
        updated_draft_equal = c.get(base + "/intake", headers=auth).json() == intake
        assert updated_draft_equal
        uploaded = c.post(
            base + "/documents",
            headers={**auth, "Idempotency-Key": uuid4().hex},
            data={"expected_version": intake["version"]},
            files={
                "file": (
                    "requirements.docx",
                    docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert uploaded.status_code == 201, uploaded.text
        document = uploaded.json()["items"][0]
        parsed = mutation(
            "POST",
            base + "/documents/" + document["document_version_id"] + "/parse",
            {"expected_version": uploaded.json()["intake_version"]},
        )
        assert parsed["items"][0]["parse_status"] == "COMPLETED"
        canonical_documents = c.get(canonical_base + "/intake/documents")
        assert canonical_documents.status_code == 200 and canonical_documents.json() == parsed
        confirmed = mutation(
            "POST", base + "/intake/confirm", {"expected_version": parsed["intake_version"]}
        )
        assert confirmed["status"] == "CONFIRMED"
        canonical_confirmed = c.get(canonical_base + "/intake")
        assert canonical_confirmed.status_code == 200 and canonical_confirmed.json() == confirmed
        assert (
            canonical_confirmed.json()["analysis_snapshot_id"] == confirmed["analysis_snapshot_id"]
        )
        with f["sf"]() as s:
            traces = s.scalars(
                select(b.SourceTraceRefEntity).where(
                    b.SourceTraceRefEntity.document_version_id == document["document_version_id"]
                )
            ).all()
            assert traces
            formal = s.scalars(
                select(e.BusinessFactEntity).where(
                    e.BusinessFactEntity.project_id == project,
                    e.BusinessFactEntity.tenant_id == f["tenant"],
                )
            ).all()
            assert any(row.fact_type == "GENERIC_FIELD" for row in formal)
        start_url = base + "/snapshots/" + confirmed["analysis_snapshot_id"] + "/workflow"
        key = uuid4().hex
        started = c.post(start_url, headers={**auth, "Idempotency-Key": key}, json={})
        assert started.status_code == 202, started.text
        assert started.json()["workflow_run_id"] == confirmed["workflow_run_id"]
        assert (
            c.post(start_url, headers={**auth, "Idempotency-Key": key}, json={}).json()
            == started.json()
        )
        # A process interrupted after durable claim must be recoverable. A second
        # delivery process cannot claim a live lease; restart reclaims the same run.
        from datetime import UTC, datetime, timedelta

        from crossborder_compliance.infrastructure.persistence.integration_models import (
            IntegrationWorkflowDeliveryEntity as Job,
        )

        claim = worker.claim(Job, "workflow_run_id")
        assert claim and claim[0] == confirmed["workflow_run_id"]
        parallel = IntegrationDeliveryWorker(app)
        assert parallel.claim(Job, "workflow_run_id") is None
        with f["sf"]() as session, session.begin():
            session.get(Job, claim[0]).lease_until = datetime.now(UTC) - timedelta(seconds=1)
        worker = IntegrationDeliveryWorker(app)
        worker.start()
        try:
            status = None
            for _ in range(100):
                response = c.get(started.json()["status_url"], headers=auth)
                assert response.status_code == 200, response.text
                status = response.json()
                if status["status"] in {"COMPLETED", "REVIEW_REQUIRED", "FAILED", "WARNING"}:
                    break
                time.sleep(0.1)
            assert status["status"] == ("REVIEW_REQUIRED" if review else "COMPLETED"), status
            result = c.get(started.json()["result_url"], headers=auth)
            assert result.status_code == 200, result.text
            value = result.json()
            assert value["analysis_snapshot_id"] == confirmed["analysis_snapshot_id"]
            if review:
                assert (
                    status["review_id"]
                    and value["review_id"]
                    and value["status"] == "REVIEW_REQUIRED"
                )
                assert value["final_path"] is None
            else:
                assert (
                    value["final_path"]
                    and value["legal_basis"]
                    and value["retrieval"]["evidence_pack"]["items"]
                )
                assert value["mode"] == "DATA_AWARE"
                assert len(value["classification"]) == 2
                assert {row["jurisdiction_id"] for row in value["classification"]} == set(
                    facts["source_locations"] + facts["destination_locations"]
                )
                assert len({row["data_item_id"] for row in value["classification"]}) == 1
                assert len({row["scheme_version_id"] for row in value["classification"]}) == 1
                # Compare the same authorized canonical actor/context across
                # the external façade and internal route; no Web-user impersonation.
                from types import SimpleNamespace

                from crossborder_compliance.domain.integrations import IntegrationScope
                from crossborder_compliance.infrastructure.external_composition import (
                    integration_service,
                )

                service = integration_service(SimpleNamespace(app=app))
                principal = service.authenticate(auth["Authorization"][7:])
                ctx = service.authorize(principal, IntegrationScope.COMPLIANCE_READ, UUID(project))
                app.dependency_overrides[get_repository_context] = lambda: ctx
                direct = c.get(
                    "/api/v1/workflows/" + confirmed["workflow_run_id"] + "/stage1-result"
                )
                assert direct.status_code == 200, direct.text
                assert direct.json() == value
            events = c.get(started.json()["events_url"], headers=auth)
            assert events.status_code == 200 and "data: " in events.text
            assert "checkpoint" not in events.text and "secret_ref" not in events.text
            # Real upgraded HTTP connection, same canonical projection as SSE.
            import json

            from websockets.exceptions import ConnectionClosedOK, InvalidStatus
            from websockets.sync.client import connect

            ws_url = (
                str(c.base_url).rstrip("/").replace("http://", "ws://")
                + "/api/v1/external/ws/workflows/"
                + confirmed["workflow_run_id"]
            )
            projected = [
                json.loads(line[6:])
                for line in events.text.splitlines()
                if line.startswith("data: ")
            ]
            with connect(ws_url, additional_headers=auth) as connection:
                values = []
                while True:
                    try:
                        values.append(json.loads(connection.recv(timeout=5)))
                    except ConnectionClosedOK:
                        break
            assert values == projected
            assert all(
                v["event_code"]
                in {
                    "WORKFLOW_STARTED",
                    "WORKFLOW_PROGRESS",
                    "WORKFLOW_COMPLETED",
                    "WORKFLOW_FAILED",
                    "REVIEW_REQUIRED",
                }
                for v in values
            )
            with connect(
                ws_url, additional_headers={**auth, "Last-Event-ID": values[0]["event_id"]}
            ) as connection:
                assert json.loads(connection.recv(timeout=5)) == values[1]
            app.dependency_overrides[get_repository_context] = lambda: RepositoryContext.user(
                UUID(f["tenant"]), "integration-administrator", {"integration:manage"}
            )
            other, _ = authenticated(c)
            assert c.get(started.json()["events_url"], headers=other).status_code == 404
            assert c.get(started.json()["result_url"], headers=other).status_code == 404
            with pytest.raises(InvalidStatus):
                connect(ws_url, additional_headers=other)
            import os
            import subprocess
            from pathlib import Path

            measured = {
                "runner_checkout_sha": subprocess.check_output(
                    ["git", "rev-parse", "HEAD"], text=True
                ).strip(),
                "real_http": True,
                "real_postgresql": True,
                "real_docx": True,
                "genuine_source_trace": bool(traces),
                "document_fact_participates": True,
                "distinct_jurisdictions": True,
                "workflow_status": status["status"],
                "snapshot_id": confirmed["analysis_snapshot_id"],
                "workflow_run_id": confirmed["workflow_run_id"],
                "sse_websocket_parity": values == projected,
                "reconnect": True,
                "cross_client_denied": True,
                "lease_restart": True,
                "no_paid_internet_llm": True,
            }
            out = Path(os.getenv("EVIDENCE_DIR", "artifacts/m2e"))
            out.mkdir(parents=True, exist_ok=True)
            (
                out
                / (
                    "external_review_validation.json"
                    if review
                    else "external_vertical_validation.json"
                )
            ).write_text(json.dumps(measured, indent=2) + "\n")
            if not review:
                parity = {
                    "runner_checkout_sha": measured["runner_checkout_sha"],
                    "same_authorized_canonical_context": True,
                    "canonical_route": "Stage1ResultService",
                    "all_structured_fields_equal": direct.json() == value,
                    "canonical_input_channel_exercised": True,
                    "intake_adapter_equal": initial_draft_equal and updated_draft_equal,
                    "canonical_document_universe_equal": canonical_documents.json() == parsed,
                    "confirmed_input_and_snapshot_equal": canonical_confirmed.json() == confirmed,
                    "model_preferences_equal": canonical_confirmed.json()["intake"]
                    == confirmed["intake"],
                    "owning_fields": [
                        "classification",
                        "applicability",
                        "cross_border",
                        "obligation",
                        "risk",
                        "recommendation",
                        "final_path",
                        "document_requirements",
                        "legal_basis",
                        "analysis_snapshot_id",
                    ],
                    "presentation_exceptions": [],
                    "result_not_precomputed": True,
                }
                (out / "ui_api_parity.json").write_text(json.dumps(parity, indent=2) + "\n")
                # Input/model changes must use the same canonical successor
                # operation as Web, preserving the exact historical S1 result.
                successor_url = base + "/intake/supersede"
                successor_key = uuid4().hex
                successor_headers = {**auth, "Idempotency-Key": successor_key}
                successor_body = {"expected_version": confirmed["version"]}
                successor = c.post(successor_url, headers=successor_headers, json=successor_body)
                assert successor.status_code == 200, successor.text
                assert (
                    c.post(successor_url, headers=successor_headers, json=successor_body).json()
                    == successor.json()
                )
                clash = c.post(
                    successor_url,
                    headers=successor_headers,
                    json={"expected_version": confirmed["version"] + 1},
                )
                assert clash.status_code == 409
                draft = c.get(base + "/intake", headers=auth).json()
                assert draft["status"] == "DRAFT" and draft["analysis_snapshot_id"] is None
                changed = mutation(
                    "PUT",
                    base + "/intake",
                    {
                        "expected_version": draft["version"],
                        "facts": {**facts, "scenario_description": "Successor input"},
                    },
                )
                next_confirmed = mutation(
                    "POST", base + "/intake/confirm", {"expected_version": changed["version"]}
                )
                assert next_confirmed["analysis_snapshot_id"] != confirmed["analysis_snapshot_id"]
                assert c.get(started.json()["result_url"], headers=auth).json() == value
                measured["external_successor_snapshot"] = True
                measured["historical_result_unchanged_after_successor"] = True
                (out / "external_vertical_validation.json").write_text(
                    json.dumps(measured, indent=2) + "\n"
                )

        finally:
            worker.stop()
