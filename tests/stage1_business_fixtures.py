"""Synthetic reviewed/published business universe; never seeded project results.

The foundation's unrelated project is configuration/bootstrap history only.
Each journey must create its own blank project through real HTTP.
"""
import copy
import io
import json
import re
from contextlib import ExitStack, contextmanager
from datetime import date
from types import SimpleNamespace
from uuid import UUID, uuid4

from cryptography.fernet import Fernet
from docx import Document
from fastapi import HTTPException
from fastapi.exception_handlers import http_exception_handler
from sqlalchemy import select
from phase1kb_http_servers import local_provider
from test_phase1l_b_postgres import fixture, foundation_i, foundation_j
from test_phase1i_postgres import publish_config
from test_m2e_vertical_postgres import server
from test_m2a_intake import setup
from test_m2e_http import http_server
from test_phase1c_model_registry import _ctx
from test_phase1k_a_configuration_postgres import publish as publish_metadata
from test_phase1kb_control_postgres import configured_model, provider_request, publish
from test_phase1h_postgres import publish as publish_rule

from crossborder_compliance.application.model_control_services import ModelControlService
from crossborder_compliance.domain.security import RepositoryContext
from crossborder_compliance.domain.llm_invocation import InvocationPurpose
from crossborder_compliance.infrastructure.llm_gateway_http import check_endpoint
from crossborder_compliance.infrastructure.llm_provider_inspection import ProviderInspection
from crossborder_compliance.infrastructure.llm_secrets import EncryptedFileSecretStore
from crossborder_compliance.infrastructure.persistence import metadata_models as m
from crossborder_compliance.infrastructure.persistence.model_control import PostgresModelControlRepository
from crossborder_compliance.infrastructure.persistence.metadata_repositories import PostgresAdminMetadataRepository
from crossborder_compliance.infrastructure.persistence.config_admin_repositories import PostgresGovernedArtifactAdminRepository
from crossborder_compliance.interfaces.api.dependencies import get_repository_context
from crossborder_compliance.infrastructure.persistence.rule_admin_repository import PostgresRuleAdminRepository
from crossborder_compliance.domain.rules import RuleContract, RuleTestCase

OUTCOMES={'DIRECT':'DIRECT_TRANSFER_ALLOWED','CONDITIONAL':'CONDITIONAL_TRANSFER_ALLOWED',
 'LOCALIZATION':'TRANSFER_NOT_ALLOWED_OR_LOCALIZATION_REQUIRED','REVIEW':'REVIEW_REQUIRED'}


def predicate(field, kind, op, value):
    return dict(entry_id=str(uuid4()),code='SYNTHETIC_GOVERNED_CONDITION',fields=[dict(code=field,field_type=dict(kind=kind))],
        predicate=json.dumps(dict(op=op,field=field,value=value)),examples=[dict(facts=[dict(code=field,value=value)],expected='TRUE')])


def configure_business(f, outcome):
    with f['sf']() as s:
        definitions={d.kind:(d.definition_id,copy.deepcopy(s.get(m.MetadataVersionEntity,d.active_version_id).payload_json))
            for d in s.scalars(select(m.MetadataDefinitionEntity).where(m.MetadataDefinitionEntity.tenant_id==f['tenant'],m.MetadataDefinitionEntity.kind.in_(
                ['OBLIGATION_POLICY','CROSS_BORDER_ASSESSMENT_POLICY','DOCUMENT_REQUIREMENT_POLICY','COMPLIANCE_PATH_POLICY']))) }
    ob_id,ob=definitions['OBLIGATION_POLICY']
    source=next(e for e in ob['entries'] if e['jurisdiction_id']==f['juri'])
    source['fulfillment_conditions']=[predicate('count','integer','gte',4 if outcome=='CONDITIONAL' else 3)]
    if outcome=='LOCALIZATION':source['legal_effect_code']='LOCALIZATION_REQUIRED'
    publish_config(f,'OBLIGATION_POLICY',ob,definition_id=ob_id)
    path_id,path=definitions['COMPLIANCE_PATH_POLICY']
    for template in path['templates']:
        template['obligation_entry_ids']=[entry['entry_id'] for entry in ob['entries']]
    publish_config(f,'COMPLIANCE_PATH_POLICY',path,definition_id=path_id)
    cb_id,cb=definitions['CROSS_BORDER_ASSESSMENT_POLICY']
    cb['transfer_legally_possible']=predicate('generic_field','string','eq','Authorized') if outcome=='REVIEW' else predicate('count','integer','gte',3)
    cb['prerequisites']=[dict(requirement_id=source['entry_id'],kind='APPROVAL')]
    publish_config(f,'CROSS_BORDER_ASSESSMENT_POLICY',cb,definition_id=cb_id)
    doc_id,doc=definitions['DOCUMENT_REQUIREMENT_POLICY']
    doc['entries']=[dict(entry_id=str(uuid4()),document_type_code='SYNTHETIC_DOCUMENT_'+level,name='Synthetic governed '+level,
        requirement_level=level,trigger_obligation_entry_ids=[source['entry_id']]) for level in ('REQUIRED','CONDITIONAL','RECOMMENDED','NOT_APPLICABLE')]
    publish_config(f,'DOCUMENT_REQUIREMENT_POLICY',doc,definition_id=doc_id)


def genuine_docx(project_id, *, conflict=False, semantic_only=False):
    d=Document();d.core_properties.identifier=project_id
    # Office identifier differs per project to respect M2-B binary ownership;
    # visible business content is the same manifest across Web/API channels.
    d.add_paragraph('Governed Data supports the service-processing purpose. STAGE1_METRIC_REQUIRES_SEMANTICS BUSINESS_DOC_MARKER')
    if conflict:d.add_paragraph('Purpose: Other purpose')
    if not semantic_only:
        table=d.add_table(rows=2,cols=1);table.cell(0,0).text='Governed Data';table.cell(1,0).text='Business measurement'
    out=io.BytesIO();d.save(out);return out.getvalue()


@contextmanager
def business_environment(tmp_path, outcome, *, provider_port=0, candidate_values=None):
    f=foundation_i.__wrapped__(fixture.__wrapped__(),SimpleNamespace(param={'no_data':True}))
    f=foundation_j.__wrapped__(f,SimpleNamespace(param={}))
    app,facts=server(f,tmp_path)
    from crossborder_compliance.interfaces.api.routes.reviews import router as review_router
    app.include_router(review_router)
    failures=[]
    @app.exception_handler(HTTPException)
    async def bounded_diagnostic(request, exc):
        cause=exc.__cause__
        reason=str(cause)
        failures.append(dict(status=exc.status_code,reason=reason if re.fullmatch('[A-Z0-9_: ]{1,100}',reason) else type(cause).__name__))
        return await http_exception_handler(request,exc)
    configure_business(f,outcome)
    _,_,base,_=setup(f)
    admin=_ctx(UUID(f['tenant']))
    metadata=PostgresAdminMetadataRepository(f['sf'],admin)
    with f['sf']() as s:
        count=s.scalar(select(m.MetadataDefinitionEntity).where(m.MetadataDefinitionEntity.tenant_id==f['tenant'],m.MetadataDefinitionEntity.kind=='BUSINESS_FACT_TYPE',m.MetadataDefinitionEntity.code=='COUNT'))
        count_id=count.definition_id
    # Bootstrap identities stay stable. Actual acceptance inputs use newly
    # reviewed/published metadata versions, not seeded ACTIVE configuration.
    for identity in (f['scenario'],f['a'],count_id):
        with f['sf']() as session:
            definition=session.get(m.MetadataDefinitionEntity,identity)
            payload=copy.deepcopy(session.get(m.MetadataVersionEntity,definition.active_version_id).payload_json)
        version=metadata.create_version(definition_id=UUID(identity),payload={**payload,'fixture_kind':'SYNTHETIC_GOVERNED_TEST_DATA'})
        published=publish_metadata(metadata,version,True)
        assert published['lifecycle_status']=='ACTIVE' and published['version_no']>=2
    extraction_type=count_id
    if outcome=='DIRECT':
        purpose=metadata.create_definition(kind='BUSINESS_FACT_TYPE',code='PROCESSING_PURPOSE',display_name='Document-derived processing purpose')
        publish_metadata(metadata,metadata.create_version(definition_id=UUID(purpose['definition_id']),payload={}),True)
        extraction_type=purpose['definition_id']
        with f['sf']() as s:
            rule=s.get(m.RuleDefinitionEntity,f['rule']['definition_id'])
            config=copy.deepcopy(s.get(m.RuleVersionEntity,rule.active_version_id).runtime_contract_json)
        config['fields']['processing_purpose']={'kind':'string'}
        config['conditions']={'op':'and','args':[config['conditions'],{'op':'eq','field':'processing_purpose','value':'service-processing'}]}
        con=RuleContract.model_validate(config)
        tests=[RuleTestCase(facts={'count':3,'processing_purpose':'service-processing'},expected_match=True,expected_actions=con.actions,expected_severity=con.severity,expected_evidence_requirement=con.evidence_required),
            RuleTestCase(facts={'count':3,'processing_purpose':'different-purpose'},expected_match=False,expected_actions=(),expected_severity=None,expected_evidence_requirement=())]
        owner=PostgresRuleAdminRepository(f['sf'],f['ctx'])
        reviewer=RepositoryContext.user(UUID(f['tenant']),'independent-reviewer',set(f['ctx'].permission.scopes))
        draft=owner.create_version(UUID(f['rule']['definition_id']),{'runtime_contract':con.model_dump(mode='json'),'tests':[t.model_dump(mode='json') for t in tests]})
        publish_rule(owner,draft,PostgresRuleAdminRepository(f['sf'],reviewer))
    prompts=PostgresGovernedArtifactAdminRepository(f['sf'],admin,'prompts')
    prompt=publish_metadata(prompts,prompts.create_draft(code=uuid4().hex,display_name='Synthetic source-addressed extraction',
        payload={'template_text':'Return source-addressed candidates only; model opinion is not legal authority.','capability_requirement':['STRUCTURED_OUTPUT']}))
    prompts.bind_llm_invocation(UUID(prompt['version_id']),InvocationPurpose.DOCUMENT_CANDIDATE_EXTRACTION)
    prompts.bind_llm_invocation(UUID(prompt['version_id']),InvocationPurpose.QUERY_EXPANSION)
    policy=metadata.create_definition(kind='LLM_INVOCATION_POLICY',code=uuid4().hex,display_name='Synthetic invocation governance')
    publish_metadata(metadata,metadata.create_version(definition_id=UUID(policy['definition_id']),payload={
        'allowed_triggers':{'STANDARD':['DOCUMENT_SEMANTIC_EXTRACTION_REQUIRED','KNOWLEDGE_EVIDENCE_INSUFFICIENT'],
            'ENHANCED':['DOCUMENT_SEMANTIC_EXTRACTION_REQUIRED','KNOWLEDGE_EVIDENCE_INSUFFICIENT']},
        'max_models':2,'candidate_fact_type_ids':[extraction_type],'candidate_min_confidence':.95}),True)
    requests=[]
    def response(payload):
        content=payload['messages'][-1]['content']
        if not content.startswith('{'):return {'ok':True}
        body=json.loads(content)
        if 'source_nodes' not in body:
            if 'original_query' in body:
                requests.append(dict(model=payload['model'],query_expansion=True))
                return {'search_phrases':['Generic', 'Unrelated scope candidate '+payload['model']]}
            return {'ok':True}
        node=next(n for n in body['source_nodes'] if 'BUSINESS_DOC_MARKER' in n['text'])
        requests.append(dict(model=payload['model'],source_node_id=node['source_node_id']))
        proposed=(candidate_values or {}).get(payload['model'],'service-processing')
        return {'facts':[{'source_node_id':node['source_node_id'],'fact_type':'PROCESSING_PURPOSE','value':proposed,'confidence':1}],
            'items':[{'source_node_id':node['source_node_id'],'name':'Governed Data','confidence':1}],'nodes':[],'edges':[]}
    with ExitStack() as stack:
        url,calls=stack.enter_context(local_provider(('A1','A2'),structured_response=response,port=provider_port))
        models=PostgresModelControlRepository(f['sf'],admin)
        secrets=EncryptedFileSecretStore(tmp_path/'private-model-secrets',Fernet.generate_key(),UUID(f['tenant']))
        control=ModelControlService(models,admin,secrets,check_endpoint)
        provider=publish(f['sf'],UUID(f['tenant']),'PROVIDER',control.save_provider(provider_request(credential='ci-local-credential').model_copy(update={'base_url':url})))
        deployments=[]
        for name in ('A1','A2'):
            model=publish(f['sf'],UUID(f['tenant']),'MODEL',configured_model(f['sf'],UUID(f['tenant']),models,provider,name))
            assert ProviderInspection(models,secrets).model_test(UUID(model['deployment_id']),'chat')['status']=='HEALTHY'
            deployments.append(model)
        definition=metadata.create_definition(kind='MODEL_USAGE_POLICY',code=uuid4().hex,display_name='Tenant business models')
        publish_metadata(metadata,metadata.create_version(definition_id=UUID(definition['definition_id']),payload={
            'scope_type':'TENANT','default_mode':'INTERNAL_MODEL_ONLY','allowed_operations':['structured_output'],
            'allowed_model_ids':[d['model_id'] for d in deployments],'allowed_provider_ids':[provider['provider_id']],
            'allowed_trust_levels':['APPROVED'],'allowed_data_boundaries':['TENANT']}),True)
        def refresh_models():
            for model in deployments:
                assert ProviderInspection(models,secrets).model_test(UUID(model["deployment_id"]),"chat")["status"]=="HEALTHY"
        def provision_project_models(project_id):
            # Configuration fixture operator publishes a genuine project policy;
            # the business actor receives no policy-admin or system bypass.
            for model in deployments:
                assert ProviderInspection(models,secrets).model_test(UUID(model['deployment_id']),'chat')['status']=='HEALTHY'
            definition=metadata.create_definition(kind='MODEL_USAGE_POLICY',code=uuid4().hex,display_name='Published business project models')
            return publish_metadata(metadata,metadata.create_version(definition_id=UUID(definition['definition_id']),payload={
                'scope_type':'PROJECT','project_id':project_id,'default_mode':'INTERNAL_MODEL_ONLY','allowed_operations':['structured_output'],
                'allowed_model_ids':[d['model_id'] for d in deployments],'allowed_provider_ids':[provider['provider_id']],
                'allowed_trust_levels':['APPROVED'],'allowed_data_boundaries':['TENANT']}),True)
        business_scopes={s for s in base.permission.scopes if s not in {'knowledge:admin','metadata:admin','metadata:review','metadata:publish'} and not s.startswith('project:')}
        business_scopes|={'project:create','project:read','project:update','project:confirm'}
        user=RepositoryContext.user(UUID(f['tenant']),'business-user',business_scopes|{
            'llm:invoke','knowledge:retrieve','document:read','document:upload','document:parse','document:unlink','workflow:review',f"prompt:{prompt['definition_id']}:use"})
        assert not user.permission.system
        app.dependency_overrides[get_repository_context]=lambda:user
        app.state.llm_secret_store_factory=lambda context:secrets
        client=stack.enter_context(http_server(app));client.timeout=180
        facts={**facts,'analysis_as_of_date':date.today().isoformat(),'data_volume':'3'}
        yield dict(**f,app=app,client=client,facts=facts,user=user,model_ids=[d['model_id'] for d in deployments],
            provider_ids=[provider['provider_id']],
            deployment_ids=[d['deployment_id'] for d in deployments],
            provider_requests=requests,provider_calls=calls,outcome=outcome,expected=OUTCOMES[outcome],metadata=metadata,prompt=prompt,
            provision_project_models=provision_project_models,refresh_models=refresh_models,failures=failures)
