from __future__ import annotations

import io
from uuid import UUID, uuid4

import pytest
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from sqlalchemy import func, select

from crossborder_compliance.application.context_services import ContextResolutionService
from crossborder_compliance.application.document_services import (
    DocumentIngestionService, DocumentParseService,
)
from crossborder_compliance.config import get_settings
from crossborder_compliance.domain.context_resolution import (
    ContextValidationStatus, DataFlowEdgeContext,
)
from crossborder_compliance.domain.document_intelligence import (
    CandidateDiagramResult, DiagramEdgeCandidate, DiagramNodeCandidate,
)
from crossborder_compliance.domain.security import RepositoryContext
from crossborder_compliance.infrastructure.document_parsers import default_native_parsers
from crossborder_compliance.infrastructure.persistence.context_models import (
    AnalysisSnapshotContextPinEntity, CandidateResolutionEntity,
    ContextConflictEntity, DataItemProductLinkDetailEntity,
    ProductContextCandidateEntity, ProductScopeResolutionEntity,
    ScenarioResolutionEntity,
)
from crossborder_compliance.infrastructure.persistence.context_repositories import (
    PostgresContextResolutionRepository,
)
from crossborder_compliance.infrastructure.persistence.db import build_session_factory
from crossborder_compliance.infrastructure.persistence.document_models import (
    BusinessFactCandidateEntity, CandidateDataFlowEdgeEntity,
    CandidateDataFlowNodeEntity, CandidateDataItemEntity,
)
from crossborder_compliance.infrastructure.persistence.document_repositories import (
    PostgresDocumentIntelligenceRepository,
)
from crossborder_compliance.infrastructure.persistence.metadata_models import (
    MetadataDefinitionEntity, MetadataVersionEntity,
)
from crossborder_compliance.infrastructure.persistence.models import (
    AnalysisSnapshotEntity, DataFlowEdgeEntity, DataItemProductLinkEntity,
    JurisdictionEntity, LegalEntityEntity, ProjectEntity, ProjectPartyEntity,
    ReviewTaskEntity, TenantEntity, WorkflowRunEntity,
)
from crossborder_compliance.infrastructure.persistence.postgres_repositories import (
    OptimisticConcurrencyError,
)


pytestmark=pytest.mark.runtime_smoke


class MemoryStorage:
    def __init__(self): self.data={}
    def put(self,*,object_key:str,content:bytes,content_type:str)->str:
        self.data[object_key]=bytes(content); return "mem://"+object_key
    def get(self,storage_ref:str)->bytes:
        return self.data[storage_ref.removeprefix("mem://")]


class CleanScan:
    def scan(self,*,content:bytes,filename:str)->bool: return True


class Queue:
    def __init__(self): self.items=[]
    def enqueue(self,*,task_id:UUID,task_type:str)->None: self.items.append((task_id,task_type))


class OCR:
    def extract(self,*,content:bytes,mime_type:str):
        return [{"page_no":1,"bbox":[1,2,100,20],"text":"Generic key: generic value","confidence":.9}]


class SelectiveSimilarity:
    def candidate_similarity(self, *, left, right):
        names={str(left.get("normalized_name") or ""),str(right.get("normalized_name") or "")}
        return .95 if names=={"type","unit"} else 0.0


class Vision:
    def analyze(self,*,content:bytes,mime_type:str):
        a,b=uuid4(),uuid4()
        return CandidateDiagramResult(
            nodes=(
                DiagramNodeCandidate(a,"System A","SYSTEM",(1,1,20,20),.95),
                DiagramNodeCandidate(b,"System B","SYSTEM",(30,1,20,20),.95),
            ),
            edges=(DiagramEdgeCandidate(uuid4(),a,b,"FLOW","LEFT_TO_RIGHT",(20,5,10,5),.9),),
            labels=("System A","System B"),confidence=.92,
        )


def _sf():
    settings=get_settings()
    assert settings.database_url.startswith("postgresql"), "real PostgreSQL is mandatory"
    return build_session_factory(settings.database_url)[1]


def _active_metadata(sf, tenant:UUID, *, kind:str, code:str, display_name:str)->UUID:
    definition_id,version_id=uuid4(),uuid4()
    with sf() as s,s.begin():
        d=MetadataDefinitionEntity(
            definition_id=str(definition_id),tenant_id=str(tenant),kind=kind,code=code,
            display_name=display_name,canonical_object_type=None,canonical_object_id=None,
            parent_definition_id=None,active_version_id=None,
        )
        s.add(d);s.flush()
        s.add(MetadataVersionEntity(
            version_id=str(version_id),definition_id=str(definition_id),tenant_id=str(tenant),
            version_no=1,lifecycle_status="ACTIVE",payload_json={"fixture":True},
            created_by="phase1e-test",approved_by="phase1e-test",
        ))
        s.flush();d.active_version_id=str(version_id)
    return definition_id


def _seed(sf, tenant:UUID, project:UUID):
    jurisdiction_id=uuid4(); legal_id=uuid4(); party_id=uuid4()
    snapshot_id=uuid4(); workflow_id=uuid4()
    with sf() as s,s.begin():
        s.add(TenantEntity(tenant_id=str(tenant),name=f"Phase1E-{tenant}"))
        s.add(ProjectEntity(
            project_id=str(project),tenant_id=str(tenant),name=f"Context-{project}",
            organization_id=None,active_version_id=None,
        ))
        s.add(JurisdictionEntity(
            jurisdiction_id=str(jurisdiction_id),tenant_id=str(tenant),
            code=f"J-{jurisdiction_id.hex[:8]}",name="Generic Jurisdiction",
            metadata_json={"canonical":True},
        ))
        s.add(LegalEntityEntity(
            legal_entity_id=str(legal_id),tenant_id=str(tenant),
            legal_name="Example Vendor Legal",registration_jurisdiction_id=str(jurisdiction_id),
        ))
        s.add(ProjectPartyEntity(
            project_party_id=str(party_id),project_id=str(project),tenant_id=str(tenant),
            legal_entity_id=str(legal_id),display_name="Example Vendor",
        ))
        s.add(AnalysisSnapshotEntity(
            analysis_snapshot_id=str(snapshot_id),tenant_id=str(tenant),
            project_version_id=str(uuid4()),snapshot_version="1.0",
            analysis_as_of_date=__import__("datetime").date.today(),
            provenance_json={"phase":"1e-test"},
        ))
        s.add(WorkflowRunEntity(
            workflow_run_id=str(workflow_id),thread_id=str(workflow_id),tenant_id=str(tenant),
            analysis_snapshot_id=str(snapshot_id),status="RUNNING",
            graph_definition_version="phase1e-test",langgraph_runtime_version="1.2.12",
            checkpointer_version="3.1.2",state_schema_version="phase1e-test",
        ))
    return jurisdiction_id,party_id,snapshot_id,workflow_id


def _pdf():
    b=io.BytesIO();c=canvas.Canvas(b)
    c.drawString(72,720,"Owner: Generic Team")
    c.drawString(72,700,"Purpose: Context resolution")
    c.showPage();c.save();return b.getvalue()


def _docx():
    from docx import Document
    d=Document();d.add_heading("Generic Requirements",1);d.add_paragraph("Owner: Generic Team")
    t=d.add_table(rows=2,cols=2)
    t.cell(0,0).text="Field";t.cell(0,1).text="Type"
    t.cell(1,0).text="temperature";t.cell(1,1).text="number"
    b=io.BytesIO();d.save(b);return b.getvalue()


def _xlsx():
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title="Fields"
    ws.append(["Field","Unit","Value"]);ws.append(["temperature","C",20])
    b=io.BytesIO();wb.save(b);return b.getvalue()


def _pptx():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches
    prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6])
    a=s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1));a.text="System A"
    z=s.shapes.add_textbox(Inches(4),Inches(1),Inches(2),Inches(1));z.text="System B"
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(3),Inches(1.5),Inches(4),Inches(1.5))
    t=s.shapes.add_table(2,2,Inches(1),Inches(3),Inches(4),Inches(1.5)).table
    t.cell(0,0).text="Field";t.cell(0,1).text="Value"
    t.cell(1,0).text="identifier";t.cell(1,1).text="sample"
    b=io.BytesIO();prs.save(b);return b.getvalue()


def _png():
    img=Image.new("RGB",(220,100),"white");d=ImageDraw.Draw(img)
    d.rectangle((10,20,80,70),outline="black");d.rectangle((140,20,210,70),outline="black")
    d.line((80,45,140,45),fill="black",width=2)
    b=io.BytesIO();img.save(b,format="PNG");return b.getvalue()


def _parse_generic_project(sf,tenant:UUID,project:UUID):
    document_repo=PostgresDocumentIntelligenceRepository(
        sf,RepositoryContext.user(tenant,"phase1e-doc")
    )
    storage=MemoryStorage();queue=Queue()
    ingest=DocumentIngestionService(document_repo,storage,CleanScan())
    parser=DocumentParseService(
        document_repo,storage,default_native_parsers(ocr=OCR(),vision=Vision()),queue
    )
    fixtures=[
        ("generic-context.pdf","application/pdf",_pdf()),
        ("generic-requirements.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document",_docx()),
        ("generic-fields.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",_xlsx()),
        ("generic-architecture.pptx","application/vnd.openxmlformats-officedocument.presentationml.presentation",_pptx()),
        ("generic-flow.png","image/png",_png()),
    ]
    for i,(filename,mime,content) in enumerate(fixtures):
        version=ingest.ingest(
            tenant_id=tenant,project_id=project,filename=filename,mime_type=mime,content=content
        )
        task=parser.request_parse(
            document_version_id=UUID(str(version["document_version_id"])),
            idempotency_key=f"phase1e-fixture-{i}",
        )
        result=parser.process_task(UUID(str(task["task_id"])))
        assert result["status"]=="COMPLETED"
    assert len(queue.items)==len(fixtures)


def test_phase1e_context_resolution_formal_inventory_flow_review_versioning_and_isolation():
    sf=_sf()
    tenant_a,tenant_b,project_a,project_b=uuid4(),uuid4(),uuid4(),uuid4()
    jurisdiction_a,party_a,snapshot_a,workflow_a=_seed(sf,tenant_a,project_a)
    _seed(sf,tenant_b,project_b)

    fact_type=_active_metadata(sf,tenant_a,kind="BUSINESS_FACT_TYPE",code="GENERIC_FIELD",display_name="Generic Field")
    product_a=_active_metadata(sf,tenant_a,kind="PRODUCT",code=f"fixture-product-a-{uuid4().hex[:6]}",display_name="Fixture Product A")
    product_b=_active_metadata(sf,tenant_a,kind="PRODUCT",code=f"fixture-product-b-{uuid4().hex[:6]}",display_name="Fixture Product B")
    scenario=_active_metadata(sf,tenant_a,kind="SCENARIO",code=f"fixture-scenario-{uuid4().hex[:6]}",display_name="Fixture Scenario")
    system_type=_active_metadata(sf,tenant_a,kind="SYSTEM_TYPE",code=f"fixture-system-{uuid4().hex[:6]}",display_name="Fixture System Type")
    device_type=_active_metadata(sf,tenant_a,kind="DEVICE_TYPE",code=f"fixture-device-{uuid4().hex[:6]}",display_name="Fixture Device Type")
    party_role=_active_metadata(sf,tenant_a,kind="PARTY_ROLE",code=f"fixture-party-{uuid4().hex[:6]}",display_name="Fixture Service Role")
    assert fact_type

    _parse_generic_project(sf,tenant_a,project_a)
    with sf() as s,s.begin():
        nodes=s.scalars(select(CandidateDataFlowNodeEntity).where(
            CandidateDataFlowNodeEntity.tenant_id==str(tenant_a)
        )).all()
        assert nodes
        for node in nodes:
            node.party_candidate="Example Vendor"
            node.location_candidate="Canonical fixture"

    repo=PostgresContextResolutionRepository(sf,RepositoryContext.user(tenant_a,"phase1e-user"))
    other=PostgresContextResolutionRepository(sf,RepositoryContext.user(tenant_b,"other-user"))

    candidate_items=repo.list_candidate_items(project_a)
    field_candidates=[x for x in candidate_items if x["normalized_name"]=="field"]
    unit_candidates=[x for x in candidate_items if x["normalized_name"]=="unit"]
    assert len(field_candidates)>=3, "DOCX + XLSX + PPTX must repeat the same generic Field"
    assert unit_candidates

    flow_edges=repo.list_candidate_flow_edges(project_a)
    valid_edge=next(x for x in flow_edges if x.get("direction")=="LEFT_TO_RIGHT")
    unknown_edge=next(x for x in flow_edges if x.get("direction") in {None,"UNKNOWN"})

    trace_id=UUID(str(field_candidates[0]["source_trace_ids"][0]))
    result=ContextResolutionService(repo,SelectiveSimilarity()).run(
        project_a,
        selected_product_scope=(product_a,),
        detected_product_scope=(product_b,),
        selected_scenarios=(scenario,),detected_scenarios=(scenario,),
        systems=(
            {"display_name":"System A","system_type_ref":str(system_type),"source_trace_ids":[str(trace_id)]},
            {"display_name":"System B","system_type_ref":str(system_type),"source_trace_ids":[str(trace_id)]},
        ),
        devices=(
            {"display_name":"Generic Device","device_type_ref":str(device_type),"system_name":"System A","source_trace_ids":[str(trace_id)]},
        ),
        parties=(
            {"display_name":"Example Vendor","role_definition_id":str(party_role),"source_trace_ids":[str(trace_id)]},
            {"display_name":"Unresolved External Party","role_definition_id":str(party_role),"source_trace_ids":[str(trace_id)]},
        ),
        jurisdictions=(
            {
                "jurisdiction_id":str(jurisdiction_a),"input_value":"Canonical fixture",
                "context_type":"PROCESSING","location_precision":"EXACT_CANONICAL",
                "source":"EXPLICIT_USER_INPUT","confidence":1.0,
                "latitude":1.0,"longitude":2.0,"source_trace_ids":[str(trace_id)],
            },
            {
                "jurisdiction_id":None,"input_value":"Unspecified region",
                "context_type":"STORAGE","location_precision":"UNKNOWN",
                "source":"DOCUMENT_DERIVED","confidence":0.4,
                "source_trace_ids":[str(trace_id)],
            },
        ),
        data_item_product_bindings={
            str(field_candidates[0]["candidate_id"]):[str(product_a)],
            str(unit_candidates[0]["candidate_id"]):[str(product_b)],
        },
        data_item_flow_bindings={
            str(valid_edge["candidate_edge_id"]):[str(field_candidates[0]["candidate_id"])],
        },
        data_groups=(
            {
                "name":"Operational Fields",
                "candidate_ids":[str(field_candidates[0]["candidate_id"]),str(unit_candidates[0]["candidate_id"])],
                "grouping_reason":"LOGICAL_CONTEXT_GROUP",
            },
        ),
        workflow_run_id=workflow_a,
    )
    assert result["version"]==1
    stats=result["statistics"]
    assert stats["raw_field_count"]>stats["normalized_data_item_count"]>0
    assert stats["data_group_count"]==1
    assert stats["data_flow_node_count"]>=2
    assert stats["data_flow_edge_count"]>=1
    assert stats["conflict_count"]>=6
    assert len(result["review_task_ids"])>=6

    business=repo.get_business_context(project_a)
    assert business
    assert all(x["validation_status"]=="REVIEW_REQUIRED" for x in business)
    assert all(x["conflict_status"]=="CONFLICT" for x in business)

    products=repo.get_product_context(project_a)
    assert products and products[0]["resolution_status"]=="REVIEW_REQUIRED"
    conflicts=repo.list_conflicts(project_a)
    conflict_types={x["conflict_type"] for x in conflicts}
    assert "PRODUCT_CONTEXT_CONFLICT" in conflict_types
    assert "BUSINESS_FACT_CONFLICT" in conflict_types
    assert "DATA_ITEM_POSSIBLE_DUPLICATE" in conflict_types
    assert "PARTY_CONTEXT_CONFLICT" in conflict_types
    assert "FLOW_DIRECTION_CONFLICT" in conflict_types
    assert "LOCATION_CONTEXT_CONFLICT" in conflict_types

    items=repo.get_data_items(project_a)
    field_item=next(x for x in items if x["canonical_name"]=="field")
    unit_item=next(x for x in items if x["canonical_name"]=="unit")
    assert len(field_item["raw_candidate_ids"])>=3
    assert len(field_item["source_trace_ids"])>=3

    field_id=UUID(str(field_item["data_item_id"]));unit_id=UUID(str(unit_item["data_item_id"]))
    with sf() as s:
        field_products=list(s.scalars(
            select(DataItemProductLinkEntity.product_ref).where(
                DataItemProductLinkEntity.tenant_id==str(tenant_a),
                DataItemProductLinkEntity.data_item_id==str(field_id),
            )
        ))
        unit_products=list(s.scalars(
            select(DataItemProductLinkEntity.product_ref).where(
                DataItemProductLinkEntity.tenant_id==str(tenant_a),
                DataItemProductLinkEntity.data_item_id==str(unit_id),
            )
        ))
    assert field_products==[str(product_a)]
    assert unit_products==[str(product_b)]

    flows=repo.get_data_flows(project_a)
    assert flows["edges"] and flows["data_item_links"]
    assert all(n["system_id"] for n in flows["nodes"])
    assert all(n["party_id"]==str(party_a) for n in flows["nodes"])
    assert all(n["jurisdiction_context_id"] for n in flows["nodes"])
    assert any(x["data_item_id"]==str(field_id) for x in flows["data_item_links"])

    jurisdictions=repo.get_jurisdiction_context(project_a)
    exact=next(x for x in jurisdictions if x["location_precision"]=="EXACT_CANONICAL")
    unknown=next(x for x in jurisdictions if x["location_precision"]=="UNKNOWN")
    assert exact["latitude"]==1.0 and exact["longitude"]==2.0
    assert unknown["latitude"] is None and unknown["longitude"] is None

    parties=repo.get_parties(project_a)
    assert any(x["project_party_id"]==str(party_a) and not x["review_required"] for x in parties)
    assert any(x["project_party_id"] is None and x["review_required"] for x in parties)

    formal_result=repo.get_context_resolution(project_a)
    assert formal_result is not None and formal_result.version==1
    assert formal_result.product_contexts
    assert formal_result.scenario_contexts
    assert len(formal_result.system_contexts)==2
    assert formal_result.device_contexts
    assert formal_result.party_contexts
    assert len(formal_result.jurisdiction_contexts)==2
    assert formal_result.conflicts
    assert formal_result.review_task_ids
    assert formal_result.data_inventory_summary["items"]
    assert formal_result.data_flow_summary["edges"]

    with sf() as s:
        product_candidate_count=int(s.scalar(select(func.count()).select_from(ProductContextCandidateEntity).where(
            ProductContextCandidateEntity.tenant_id==str(tenant_a),
            ProductContextCandidateEntity.project_id==str(project_a),
            ProductContextCandidateEntity.version==1,
        )) or 0)
        scenario_resolution_count=int(s.scalar(select(func.count()).select_from(ScenarioResolutionEntity).where(
            ScenarioResolutionEntity.tenant_id==str(tenant_a),
            ScenarioResolutionEntity.project_id==str(project_a),
            ScenarioResolutionEntity.version==1,
        )) or 0)
        assert product_candidate_count==2
        assert scenario_resolution_count==2

    with sf() as s:
        counts={
            "fact_candidates":int(s.scalar(select(func.count()).select_from(BusinessFactCandidateEntity).where(BusinessFactCandidateEntity.tenant_id==str(tenant_a))) or 0),
            "item_candidates":int(s.scalar(select(func.count()).select_from(CandidateDataItemEntity).where(CandidateDataItemEntity.tenant_id==str(tenant_a))) or 0),
            "node_candidates":int(s.scalar(select(func.count()).select_from(CandidateDataFlowNodeEntity).where(CandidateDataFlowNodeEntity.tenant_id==str(tenant_a))) or 0),
            "edge_candidates":int(s.scalar(select(func.count()).select_from(CandidateDataFlowEdgeEntity).where(CandidateDataFlowEdgeEntity.tenant_id==str(tenant_a))) or 0),
        }
        resolutions={
            kind:int(s.scalar(select(func.count()).select_from(CandidateResolutionEntity).where(
                CandidateResolutionEntity.tenant_id==str(tenant_a),
                CandidateResolutionEntity.candidate_type==kind,
                CandidateResolutionEntity.version==1,
            )) or 0)
            for kind in ["BUSINESS_FACT","DATA_ITEM","DATA_FLOW_NODE","DATA_FLOW_EDGE"]
        }
        review_count=int(s.scalar(select(func.count()).select_from(ReviewTaskEntity).where(
            ReviewTaskEntity.tenant_id==str(tenant_a),
            ReviewTaskEntity.review_type=="CONTEXT_RESOLUTION_REVIEW",
        )) or 0)
    assert resolutions["fact_candidates" if False else "BUSINESS_FACT"]==counts["fact_candidates"]
    assert resolutions["DATA_ITEM"]==counts["item_candidates"]
    assert resolutions["DATA_FLOW_NODE"]==counts["node_candidates"]
    assert resolutions["DATA_FLOW_EDGE"]==counts["edge_candidates"]
    assert review_count>=6

    product_conflict=next(x for x in conflicts if x["conflict_type"]=="PRODUCT_CONTEXT_CONFLICT")
    resolved=repo.resolve_conflict(
        UUID(str(product_conflict["conflict_id"])),
        resolution={"effective_product_scope":[str(product_a)]},
        resolved_by="human-reviewer",expected_record_version=1,
    )
    assert resolved["resolution_status"]=="RESOLVED"
    with sf() as s:
        scope=s.scalar(select(ProductScopeResolutionEntity).where(
            ProductScopeResolutionEntity.tenant_id==str(tenant_a),
            ProductScopeResolutionEntity.conflict_id==str(product_conflict["conflict_id"]),
        ))
        assert scope.effective_product_scope_json==[str(product_a)]
        assert scope.review_required is False
    with pytest.raises(OptimisticConcurrencyError):
        repo.resolve_conflict(
            UUID(str(product_conflict["conflict_id"])),
            resolution={"effective_product_scope":[str(product_a)]},
            resolved_by="stale-reviewer",expected_record_version=1,
        )

    run1=UUID(str(result["context_resolution_run_id"]))
    repo.pin_snapshot_context(
        analysis_snapshot_id=snapshot_a,project_id=project_a,context_resolution_run_id=run1
    )

    result2=ContextResolutionService(repo).run(
        project_a,
        selected_product_scope=(product_a,),detected_product_scope=(product_a,),
        selected_scenarios=(scenario,),detected_scenarios=(scenario,),
        systems=(
            {"display_name":"System A","system_type_ref":str(system_type),"source_trace_ids":[str(trace_id)]},
            {"display_name":"System B","system_type_ref":str(system_type),"source_trace_ids":[str(trace_id)]},
        ),
        data_item_flow_bindings={
            str(valid_edge["candidate_edge_id"]):[str(field_candidates[0]["candidate_id"])],
        },
        workflow_run_id=workflow_a,
    )
    assert result2["version"]==2
    with pytest.raises(ValueError):
        repo.pin_snapshot_context(
            analysis_snapshot_id=snapshot_a,project_id=project_a,
            context_resolution_run_id=UUID(str(result2["context_resolution_run_id"])),
        )
    with sf() as s:
        pin=s.scalar(select(AnalysisSnapshotContextPinEntity).where(
            AnalysisSnapshotContextPinEntity.tenant_id==str(tenant_a),
            AnalysisSnapshotContextPinEntity.analysis_snapshot_id==str(snapshot_a),
        ))
        assert pin.context_resolution_run_id==str(run1)
        assert pin.context_resolution_version==1

    before=len(repo.get_data_flows(project_a)["edges"])
    source_node=UUID(str(repo.get_data_flows(project_a)["nodes"][0]["flow_node_id"]))
    with pytest.raises(ValueError):
        repo.create_formal_flow_edge(
            project_id=project_a,source_node_id=source_node,target_node_id=source_node,
            flow_type="INVALID_SELF",
            detail=DataFlowEdgeContext(
                uuid4(),None,"SELF",{}, {},(trace_id,),1.0,
                ContextValidationStatus.VALIDATED,2,
            ),
        )
    assert len(repo.get_data_flows(project_a)["edges"])==before

    assert other.get_business_context(project_a)==[]
    assert other.get_product_context(project_a)==[]
    assert other.get_data_items(project_a)==[]
    assert other.get_data_flows(project_a)=={"nodes":[],"edges":[],"data_item_links":[]}
    assert other.get_context_resolution(project_a) is None


def test_phase1e_noncanonical_coordinates_are_rejected():
    sf=_sf();tenant,project=uuid4(),uuid4();_seed(sf,tenant,project)
    repo=PostgresContextResolutionRepository(sf,RepositoryContext.user(tenant,"coordinate-test"))
    from crossborder_compliance.application.context_services import JurisdictionResolutionService
    from crossborder_compliance.domain.context_resolution import LocationPrecision
    with pytest.raises(ValueError):
        JurisdictionResolutionService(repo).resolve(
            project,jurisdiction_id=None,input_value="free text place",
            context_type="PROCESSING",precision=LocationPrecision.UNKNOWN,
            source="DOCUMENT_DERIVED",confidence=.5,
            latitude=12.3,longitude=45.6,version=1,
        )
