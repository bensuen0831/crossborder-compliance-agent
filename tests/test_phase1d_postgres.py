from __future__ import annotations

import io
from dataclasses import dataclass
from uuid import UUID, uuid4

import pytest
from PIL import Image, ImageDraw
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from sqlalchemy import select

from crossborder_compliance.application.document_services import (
    CrossDocumentLinkingService, DocumentAnalysisService, DocumentIngestionService, DocumentParseService
)
from crossborder_compliance.config import get_settings
from crossborder_compliance.domain.document_intelligence import (
    CandidateDiagramResult, DiagramEdgeCandidate, DiagramNodeCandidate
)
from crossborder_compliance.domain.security import RepositoryContext
from crossborder_compliance.infrastructure.document_parsers import (
    PDFParserAdapter, TextParserAdapter, default_native_parsers
)
from crossborder_compliance.infrastructure.persistence.db import build_session_factory
from crossborder_compliance.infrastructure.persistence.document_models import (
    AnalysisSnapshotParseRunPinEntity, CandidateDataItemEntity, CandidateDataItemSourceLinkEntity, CrossDocumentLinkEntity
)
from crossborder_compliance.infrastructure.persistence.document_repositories import (
    PostgresDocumentIntelligenceRepository
)
from crossborder_compliance.infrastructure.persistence.models import (
    AnalysisSnapshotEntity, ProjectEntity, TenantEntity
)


class MemoryStorage:
    def __init__(self): self.data={}
    def put(self,*,object_key:str,content:bytes,content_type:str)->str:
        self.data[object_key]=bytes(content); return "mem://"+object_key
    def get(self,storage_ref:str)->bytes:
        return self.data[storage_ref.removeprefix("mem://")]


class CleanScan:
    def scan(self,*,content:bytes,filename:str)->bool: return True


class Queue:
    def __init__(self): self.items=[]
    def enqueue(self,*,task_id:UUID,task_type:str)->None: self.items.append((task_id,task_type))


class OCR:
    def __init__(self,confidence=.35): self.confidence=confidence
    def extract(self,*,content:bytes,mime_type:str):
        return [{"page_no":1,"bbox":[1,2,100,20],"text":"Scanned key: scanned value","confidence":self.confidence}]


class Vision:
    def analyze(self,*,content:bytes,mime_type:str):
        a,b=uuid4(),uuid4()
        return CandidateDiagramResult(
            nodes=(DiagramNodeCandidate(a,"System A","SYSTEM",(1,1,20,20),.9),DiagramNodeCandidate(b,"System B","SYSTEM",(30,1,20,20),.9)),
            edges=(DiagramEdgeCandidate(uuid4(),a,b,"FLOW","LEFT_TO_RIGHT",(20,5,10,5),.8),),
            confidence=.85,
        )


class FailingTextParser:
    parser_name="failing-test-parser"; parser_version="1"
    def supports(self,*,mime_type:str,filename:str)->bool: return mime_type=="text/plain"
    def parse(self,**kwargs): raise RuntimeError("synthetic parser failure")


def _sf():
    settings=get_settings()
    if not settings.database_url.startswith("postgresql"):
        pytest.skip("Phase 1D integration requires PostgreSQL")
    return build_session_factory(settings.database_url)[1]


def _seed(sf,tenant:UUID,project:UUID):
    with sf() as s,s.begin():
        s.add(TenantEntity(tenant_id=str(tenant),name=f"T-{tenant}"))
        s.add(ProjectEntity(project_id=str(project),tenant_id=str(tenant),name=f"P-{project}",organization_id=None,active_version_id=None))


def _text_pdf():
    b=io.BytesIO(); c=canvas.Canvas(b); c.drawString(72,720,"Owner: Generic Team"); c.drawString(72,700,"Purpose: Document foundation"); c.showPage(); c.save(); return b.getvalue()


def _scanned_pdf():
    img=Image.new("RGB",(240,100),"white"); d=ImageDraw.Draw(img); d.text((10,35),"scan",fill="black")
    ib=io.BytesIO(); img.save(ib,format="PNG"); ib.seek(0)
    b=io.BytesIO(); c=canvas.Canvas(b); c.drawImage(ImageReader(ib),50,620,width=240,height=100); c.showPage(); c.save(); return b.getvalue()


def _xlsx():
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title="Fields"; ws.append(["Field","Unit","Value"]); ws.append(["temperature","C",20]); ws.append(["pressure","kPa",101])
    b=io.BytesIO(); wb.save(b); return b.getvalue()


def _pptx():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches
    prs=Presentation(); s=prs.slides.add_slide(prs.slide_layouts[6])
    a=s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1)); a.text="System A"
    z=s.shapes.add_textbox(Inches(4),Inches(1),Inches(2),Inches(1)); z.text="System B"
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(3),Inches(1.5),Inches(4),Inches(1.5))
    t=s.shapes.add_table(2,2,Inches(1),Inches(3),Inches(4),Inches(1.5)).table
    t.cell(0,0).text="Field";t.cell(0,1).text="Value";t.cell(1,0).text="identifier";t.cell(1,1).text="x"
    b=io.BytesIO();prs.save(b);return b.getvalue()


@pytest.mark.runtime_smoke
def test_phase1d_multiformat_pipeline_provenance_summary_tenant_and_snapshot_pin():
    sf=_sf(); ta,tb,pa,pb=uuid4(),uuid4(),uuid4(),uuid4(); _seed(sf,ta,pa);_seed(sf,tb,pb)
    ra=PostgresDocumentIntelligenceRepository(sf,RepositoryContext.user(ta,"user-a"))
    rb=PostgresDocumentIntelligenceRepository(sf,RepositoryContext.user(tb,"user-b"))
    storage=MemoryStorage(); queue=Queue(); ingest=DocumentIngestionService(ra,storage,CleanScan())
    parse=DocumentParseService(ra,storage,default_native_parsers(ocr=OCR(.9),vision=Vision()),queue)

    versions=[]
    for filename,mime,data in [
        ("generic.pdf","application/pdf",_text_pdf()),
        ("generic-fields.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",_xlsx()),
        ("generic-architecture.pptx","application/vnd.openxmlformats-officedocument.presentationml.presentation",_pptx()),
    ]:
        versions.append(ingest.ingest(tenant_id=ta,project_id=pa,filename=filename,mime_type=mime,content=data))

    runs=[]
    for v in versions:
        vid=UUID(str(v["document_version_id"]))
        first=parse.request_parse(document_version_id=vid,idempotency_key="same-request")
        duplicate=parse.request_parse(document_version_id=vid,idempotency_key="same-request")
        assert first["task_id"]==duplicate["task_id"]
        runs.append(parse.process_task(UUID(str(first["task_id"]))))
    assert len(queue.items)==3
    assert all(r["status"]=="COMPLETED" for r in runs)

    # Cross-tenant UUID guessing cannot read version/run/structure.
    assert rb.get_document_version(UUID(str(versions[0]["document_version_id"]))) is None
    assert rb.get_parse_run(UUID(str(runs[0]["parse_run_id"]))) is None
    assert rb.list_structure(UUID(str(runs[0]["parse_run_id"]))) == []

    # Candidate spreadsheet item traces to exact source row/column.
    xrun=UUID(str(runs[1]["parse_run_id"]))
    items=ra.list_candidate_items(xrun)
    assert {i["normalized_name"] for i in items}>={"field","unit","value"}
    with sf() as s:
        xlsx_field=s.scalar(select(CandidateDataItemEntity).where(
            CandidateDataItemEntity.tenant_id==str(ta),
            CandidateDataItemEntity.parse_run_id==str(xrun),
            CandidateDataItemEntity.normalized_name=="field",
        ))
        assert xlsx_field is not None
        link=s.scalar(select(CandidateDataItemSourceLinkEntity).where(
            CandidateDataItemSourceLinkEntity.tenant_id==str(ta),
            CandidateDataItemSourceLinkEntity.candidate_data_item_id==xlsx_field.candidate_data_item_id,
        ))
        assert link is not None
    trace=ra.get_source_trace(UUID(str(link.source_trace_ref_id)))
    assert trace is not None
    assert trace.locator.row_index==1 and trace.locator.column_index is not None and trace.locator.sheet_name=="Fields"
    assert trace.original_text_hash
    assert rb.get_source_trace(trace.source_trace_ref_id) is None

    # Cross-document linking is based on normalized candidate + evidence, never filename.
    prun=UUID(str(runs[2]["parse_run_id"]))
    with sf() as s:
        left=s.scalar(select(CandidateDataItemEntity).where(CandidateDataItemEntity.tenant_id==str(ta),CandidateDataItemEntity.parse_run_id==str(xrun),CandidateDataItemEntity.normalized_name=="field"))
        right=s.scalar(select(CandidateDataItemEntity).where(CandidateDataItemEntity.tenant_id==str(ta),CandidateDataItemEntity.parse_run_id==str(prun),CandidateDataItemEntity.normalized_name=="field"))
        llink=s.scalar(select(CandidateDataItemSourceLinkEntity).where(CandidateDataItemSourceLinkEntity.candidate_data_item_id==left.candidate_data_item_id))
        rlink=s.scalar(select(CandidateDataItemSourceLinkEntity).where(CandidateDataItemSourceLinkEntity.candidate_data_item_id==right.candidate_data_item_id))
    lt=ra.get_source_trace(UUID(str(llink.source_trace_ref_id)))
    rt=ra.get_source_trace(UUID(str(rlink.source_trace_ref_id)))
    assert lt is not None and rt is not None
    cross=CrossDocumentLinkingService(ra).link_same_normalized_candidate(project_id=pa,left_object_id=UUID(left.candidate_data_item_id),right_object_id=UUID(right.candidate_data_item_id),left_normalized_name=left.normalized_name,right_normalized_name=right.normalized_name,left_trace=lt,right_trace=rt)
    assert cross is not None
    with sf() as s:
        assert s.get(CrossDocumentLinkEntity,str(cross.cross_document_link_id)) is not None

    # PPT connector becomes candidate flow only.
    flows=ra.list_candidate_flows(UUID(str(runs[2]["parse_run_id"])))
    assert len(flows["nodes"])>=2 and len(flows["edges"])>=1

    summary=DocumentAnalysisService(ra).summary(pa)
    assert summary.document_count==3 and summary.document_version_count==3
    assert summary.page_count>=1 and summary.sheet_count>=1 and summary.slide_count>=1
    assert summary.table_count>=2
    assert summary.raw_field_count>=summary.normalized_candidate_data_item_count>=3
    assert summary.candidate_data_flow_count>=1

    # Snapshot pins exact parse run and cannot silently switch on reparse.
    snapshot=uuid4()
    with sf() as s,s.begin():
        s.add(AnalysisSnapshotEntity(analysis_snapshot_id=str(snapshot),tenant_id=str(ta),project_version_id=str(uuid4()),snapshot_version="1.0",analysis_as_of_date=__import__("datetime").date.today(),provenance_json={"phase":"1d"}))
    docv=UUID(str(versions[0]["document_version_id"])); oldrun=UUID(str(runs[0]["parse_run_id"]))
    ra.pin_parse_run(analysis_snapshot_id=snapshot,document_version_id=docv,parse_run_id=oldrun)
    task2=parse.request_parse(document_version_id=docv,idempotency_key="reparse-2"); newrun=parse.process_task(UUID(str(task2["task_id"])))
    assert newrun["parse_run_version"]==2
    with pytest.raises(ValueError):
        ra.pin_parse_run(analysis_snapshot_id=snapshot,document_version_id=docv,parse_run_id=UUID(str(newrun["parse_run_id"])))
    with sf() as s:
        pin=s.scalar(select(AnalysisSnapshotParseRunPinEntity).where(AnalysisSnapshotParseRunPinEntity.analysis_snapshot_id==str(snapshot)))
        assert pin.parse_run_id==str(oldrun)


@pytest.mark.runtime_smoke
def test_phase1d_async_failed_retry_quality_gate_and_no_binary_db_source():
    sf=_sf(); tenant,project=uuid4(),uuid4(); _seed(sf,tenant,project)
    repo=PostgresDocumentIntelligenceRepository(sf,RepositoryContext.user(tenant,"user"))
    storage=MemoryStorage();queue=Queue();ingest=DocumentIngestionService(repo,storage,CleanScan())

    txt=ingest.ingest(tenant_id=tenant,project_id=project,filename="failure.txt",mime_type="text/plain",content=b"Key: Value")
    failing=DocumentParseService(repo,storage,[FailingTextParser()],queue)
    task=failing.request_parse(document_version_id=UUID(str(txt["document_version_id"])),idempotency_key="retry-key")
    with pytest.raises(RuntimeError): failing.process_task(UUID(str(task["task_id"])))
    assert repo.get_parse_task(UUID(str(task["task_id"])))["status"]=="FAILED"
    retry=failing.request_parse(document_version_id=UUID(str(txt["document_version_id"])),idempotency_key="retry-key")
    assert retry["task_id"]==task["task_id"] and retry["attempts"]==1 and retry["status"]=="ACCEPTED"

    # Same persisted task can be completed by a healthy worker adapter.
    healthy=DocumentParseService(repo,storage,[TextParserAdapter()],queue)
    result=healthy.process_task(UUID(str(task["task_id"])))
    assert result["status"]=="COMPLETED"

    scan=ingest.ingest(tenant_id=tenant,project_id=project,filename="scan.pdf",mime_type="application/pdf",content=_scanned_pdf())
    ocr_service=DocumentParseService(repo,storage,[PDFParserAdapter(OCR(.35))],queue)
    t=ocr_service.request_parse(document_version_id=UUID(str(scan["document_version_id"])),idempotency_key="scan")
    run=ocr_service.process_task(UUID(str(t["task_id"])))
    q=repo.get_quality(UUID(str(run["parse_run_id"])))
    assert run["ocr_used"] is True and q["status"]=="REVIEW_REQUIRED" and q["ocr_confidence"]<0.6
