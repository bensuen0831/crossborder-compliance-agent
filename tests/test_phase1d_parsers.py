from __future__ import annotations
import io
from uuid import uuid4
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

from crossborder_compliance.domain.document_intelligence import (
    CandidateDiagramResult, DiagramNodeCandidate, DiagramEdgeCandidate, CanonicalNodeType
)
from crossborder_compliance.infrastructure.document_parsers import (
    PDFParserAdapter,DOCXParserAdapter,XLSXParserAdapter,CSVParserAdapter,PPTXParserAdapter,TextParserAdapter,ImageParserAdapter
)

class FakeOCR:
    def extract(self,*,content:bytes,mime_type:str):
        return [{"page_no":1,"bbox":[10,20,200,40],"text":"Scanned field: value","confidence":0.42}]

class FakeVision:
    def analyze(self,*,content:bytes,mime_type:str):
        a,b=uuid4(),uuid4()
        return CandidateDiagramResult(
            nodes=(DiagramNodeCandidate(a,"System A","SYSTEM",(1,1,20,20),0.9),DiagramNodeCandidate(b,"System B","SYSTEM",(30,1,20,20),0.9)),
            edges=(DiagramEdgeCandidate(uuid4(),a,b,"FLOW","LEFT_TO_RIGHT",(20,5,10,5),0.8),),
            labels=("System A","System B"),confidence=0.85,
        )

def _pdf_text_bytes():
    buf=io.BytesIO(); c=canvas.Canvas(buf); c.drawString(72,720,"Generic Requirement: retain source trace"); c.showPage(); c.save(); return buf.getvalue()

def _pdf_scanned_bytes():
    img=Image.new("RGB",(300,120),"white"); d=ImageDraw.Draw(img); d.text((20,40),"Scanned field",fill="black")
    ib=io.BytesIO(); img.save(ib,format="PNG"); ib.seek(0)
    buf=io.BytesIO(); c=canvas.Canvas(buf); c.drawImage(ImageReader(ib),50,600,width=300,height=120); c.showPage(); c.save(); return buf.getvalue()

def _docx_bytes():
    from docx import Document
    d=Document(); d.add_heading("Generic Requirements",1); d.add_paragraph("Owner: Example Team")
    t=d.add_table(rows=2,cols=2); t.cell(0,0).text="Field";t.cell(0,1).text="Type";t.cell(1,0).text="measurement";t.cell(1,1).text="number"
    b=io.BytesIO(); d.save(b); return b.getvalue()

def _xlsx_bytes():
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title="Fields"; ws.append(["Field","Unit","Value"]); ws.append(["temperature","C","=1+1"]); ws.merge_cells("A4:B4"); ws["A4"]="Notes"
    ws2=wb.create_sheet("Second"); ws2.append(["Field","Value"]); ws2.append(["pressure",101])
    b=io.BytesIO(); wb.save(b); return b.getvalue()

def _pptx_bytes():
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_CONNECTOR
    prs=Presentation(); s=prs.slides.add_slide(prs.slide_layouts[6])
    a=s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1)); a.text="System A"
    b=s.shapes.add_textbox(Inches(4),Inches(1),Inches(2),Inches(1)); b.text="System B"
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(3),Inches(1.5),Inches(4),Inches(1.5))
    table=s.shapes.add_table(2,2,Inches(1),Inches(3),Inches(4),Inches(1.5)).table
    table.cell(0,0).text="Field";table.cell(0,1).text="Value";table.cell(1,0).text="identifier";table.cell(1,1).text="sample"
    out=io.BytesIO(); prs.save(out); return out.getvalue()

def _image_bytes():
    img=Image.new("RGB",(200,100),"white"); d=ImageDraw.Draw(img); d.rectangle((10,20,70,70),outline="black");d.rectangle((120,20,190,70),outline="black");d.line((70,45,120,45),fill="black",width=2)
    b=io.BytesIO();img.save(b,format="PNG");return b.getvalue()

def test_generic_pdf_native_and_scanned_ocr_fallback():
    vid=uuid4()
    native=PDFParserAdapter(FakeOCR()).parse(document_version_id=vid,content=_pdf_text_bytes(),filename="generic.pdf",mime_type="application/pdf")
    assert native.native_parse_used and not native.ocr_used
    assert any(n.node_type==CanonicalNodeType.PARAGRAPH and "Generic Requirement" in (n.original_text or "") for n in native.nodes)
    scanned=PDFParserAdapter(FakeOCR()).parse(document_version_id=vid,content=_pdf_scanned_bytes(),filename="generic-scanned.pdf",mime_type="application/pdf")
    assert scanned.ocr_used
    ocr_nodes=[n for n in scanned.nodes if "ocr_confidence" in n.metadata]
    assert ocr_nodes and ocr_nodes[0].bbox==(10,20,200,40)

def test_generic_docx_paragraph_heading_and_table_provenance():
    parsed=DOCXParserAdapter().parse(document_version_id=uuid4(),content=_docx_bytes(),filename="generic-requirements.docx",mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert any(n.node_type==CanonicalNodeType.HEADING for n in parsed.nodes)
    cells=[n for n in parsed.nodes if n.node_type==CanonicalNodeType.TABLE_CELL]
    assert cells and {"row_index","column_index","table_id"}<=set(cells[0].metadata)

def test_generic_xlsx_multi_sheet_headers_formula_merged_and_cell_provenance():
    parsed=XLSXParserAdapter().parse(document_version_id=uuid4(),content=_xlsx_bytes(),filename="generic-fields.xlsx",mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert {n.sheet_name for n in parsed.nodes if n.node_type==CanonicalNodeType.SHEET}=={"Fields","Second"}
    cells=[n for n in parsed.nodes if n.node_type==CanonicalNodeType.TABLE_CELL]
    assert any(n.metadata.get("header") and n.metadata["row_index"]==1 for n in cells)
    assert any(n.metadata.get("formula") for n in cells)
    assert any(n.metadata.get("merged_range") for n in cells)
    assert all("row_index" in n.metadata and "column_index" in n.metadata for n in cells)

def test_generic_csv_structured_not_plain_text():
    parsed=CSVParserAdapter().parse(document_version_id=uuid4(),content=b"Field,Value\nalpha,1\nbeta,2\n",filename="generic.csv",mime_type="text/csv")
    assert sum(n.node_type==CanonicalNodeType.TABLE_CELL for n in parsed.nodes)==6
    assert any(n.metadata.get("header") for n in parsed.nodes)

def test_generic_pptx_slide_table_and_diagram_candidate():
    parsed=PPTXParserAdapter().parse(document_version_id=uuid4(),content=_pptx_bytes(),filename="generic-architecture.pptx",mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert any(n.node_type==CanonicalNodeType.SLIDE for n in parsed.nodes)
    assert any(n.node_type==CanonicalNodeType.TABLE for n in parsed.nodes)
    diagrams=[n for n in parsed.nodes if n.node_type==CanonicalNodeType.DIAGRAM]
    assert diagrams and diagrams[0].metadata["candidate_edges"]

def test_generic_text_and_image_vision_candidate_have_no_legal_result():
    txt=TextParserAdapter().parse(document_version_id=uuid4(),content=b"Owner: Example\nPurpose: Testing",filename="generic.txt",mime_type="text/plain")
    assert sum(n.node_type==CanonicalNodeType.PARAGRAPH for n in txt.nodes)==2
    img=ImageParserAdapter(FakeVision()).parse(document_version_id=uuid4(),content=_image_bytes(),filename="generic-flow.png",mime_type="image/png")
    d=next(n for n in img.nodes if n.node_type==CanonicalNodeType.DIAGRAM)
    assert d.metadata["candidate_nodes"] and d.metadata["candidate_edges"]
    assert not any(k in d.metadata for k in ["legal_conclusion","classification","compliance_path","risk_score"])
