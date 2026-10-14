"""Native bounded XLSX/CSV parser adapter; never publishes catalogue proposals."""

import csv
import io
import zipfile

from openpyxl import load_workbook

from crossborder_compliance.application.knowledge_manifest import (
    ManifestError,
    ManifestPolicy,
    normalize_relative_path,
)


class NativeManifestParser:
    def rows(self, content: bytes, kind: str, policy: ManifestPolicy):
        if len(content) > policy.max_bytes:
            raise ManifestError("MANIFEST_SIZE_EXCEEDED")
        if kind == "CSV":
            try:
                yield from (
                    (row, set())
                    for row in csv.reader(io.StringIO(content.decode(policy.csv_encoding)))
                )
            except (UnicodeError, csv.Error) as exc:
                raise ManifestError("INVALID_CSV") from exc
            return
        if kind != "XLSX":
            raise ManifestError("UNSUPPORTED_MANIFEST_FORMAT")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > policy.max_entries
                    or sum(e.file_size for e in entries) > policy.max_expanded_bytes
                ):
                    raise ManifestError("MANIFEST_EXPANSION_LIMIT")
                if len({e.filename for e in entries}) != len(entries):
                    raise ManifestError("DUPLICATE_WORKBOOK_ENTRY")
                for entry in entries:
                    normalize_relative_path(entry.filename.rstrip("/"))
                    if (
                        entry.filename.startswith("xl/externalLinks/")
                        or "vbaProject" in entry.filename
                    ):
                        raise ManifestError("UNSAFE_WORKBOOK_CONTENT")
                    if entry.filename.endswith((".xml", ".rels")):
                        xml = archive.read(entry).upper()
                        if b"\x00" in xml or b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                            raise ManifestError("UNSAFE_WORKBOOK_XML")
            book = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                if policy.sheet_name is None and len(book.sheetnames) != 1:
                    raise ManifestError("SHEET_SELECTION_REQUIRED")
                sheet = book[policy.sheet_name] if policy.sheet_name else book.active
                if sheet.max_row is not None and sheet.max_row > policy.max_rows + 1:
                    raise ManifestError("MANIFEST_ROW_LIMIT")
                if sheet.max_column is not None and sheet.max_column > policy.max_columns:
                    raise ManifestError("MANIFEST_COLUMN_LIMIT")
                for row in sheet.iter_rows():
                    yield (
                        [c.value for c in row],
                        {i for i, c in enumerate(row) if c.data_type == "f"},
                    )
            finally:
                book.close()
        except ManifestError:
            raise
        except Exception as exc:
            raise ManifestError("INVALID_XLSX") from exc
