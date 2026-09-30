from __future__ import annotations

import csv
import io
import re
from pathlib import PurePath
from uuid import UUID, uuid4

from crossborder_compliance.application.document_ports import OCRPort, VisionDiagramPort
from crossborder_compliance.domain.document_intelligence import (
    CandidateDiagramResult,
    CanonicalNodeType,
    CanonicalStructureNode,
    ParsedDocument,
    new_structure_node,
)


def _norm(text: object) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _node(document_version_id: UUID, node_type: CanonicalNodeType, sequence: int, **kwargs):
    return new_structure_node(document_version_id=document_version_id, node_type=node_type, sequence=sequence, **kwargs)


class PDFParserAdapter:
    parser_name = "pypdf"
    parser_version = "1.0"

    def __init__(self, ocr: OCRPort | None = None):
        self.ocr = ocr

    def supports(self, *, mime_type: str, filename: str) -> bool:
        return mime_type == "application/pdf" or filename.lower().endswith(".pdf")

    def parse(self, *, document_version_id: UUID, content: bytes, filename: str, mime_type: str) -> ParsedDocument:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        nodes=[]; seq=0; total_text=0
        root=_node(document_version_id,CanonicalNodeType.DOCUMENT,seq,original_text=filename,normalized_text=filename,parser_confidence=1.0); nodes.append(root); seq+=1
        for page_no,page in enumerate(reader.pages,start=1):
            p=_node(document_version_id,CanonicalNodeType.PAGE,seq,parent_node_id=root.structure_node_id,page_no=page_no,parser_confidence=1.0); nodes.append(p); seq+=1
            text=page.extract_text() or ""; total_text+=len(text.strip())
            for para in [x.strip() for x in re.split(r"\n\s*\n|\n",text) if x.strip()]:
                kind=CanonicalNodeType.HEADING if len(para)<120 and (para.isupper() or para.endswith(":")) else CanonicalNodeType.PARAGRAPH
                nodes.append(_node(document_version_id,kind,seq,parent_node_id=p.structure_node_id,page_no=page_no,original_text=para,normalized_text=_norm(para),parser_confidence=0.95,source_locator={"page_no":page_no},provenance={"parser":self.parser_name})); seq+=1
        if total_text>0:
            return ParsedDocument(tuple(nodes),native_parse_used=True,detected_language="und")
        if self.ocr is None:
            return ParsedDocument(tuple(nodes),native_parse_used=True,warnings=("native PDF text extraction empty; OCR adapter unavailable",),detected_language=None)
        ocr_rows=self.ocr.extract(content=content,mime_type=mime_type)
        for row in ocr_rows:
            text=_norm(row.get("text"))
            if not text: continue
            page_no=int(row.get("page_no") or 1)
            nodes.append(_node(document_version_id,CanonicalNodeType.PARAGRAPH,seq,page_no=page_no,bbox=tuple(row["bbox"]) if row.get("bbox") else None,original_text=text,normalized_text=text,parser_confidence=float(row.get("confidence",0.5)),source_locator={"page_no":page_no,"bbox":row.get("bbox")},metadata={"ocr_confidence":float(row.get("confidence",0.5))},provenance={"parser":"ocr-fallback"})); seq+=1
        return ParsedDocument(tuple(nodes),native_parse_used=True,ocr_used=True,warnings=("OCR fallback used",),detected_language="und")


class DOCXParserAdapter:
    parser_name="python-docx"
    parser_version="1.0"
    def supports(self,*,mime_type:str,filename:str)->bool:
        return mime_type=="application/vnd.openxmlformats-officedocument.wordprocessingml.document" or filename.lower().endswith(".docx")
    def parse(self,*,document_version_id:UUID,content:bytes,filename:str,mime_type:str)->ParsedDocument:
        from docx import Document
        doc=Document(io.BytesIO(content)); nodes=[]; seq=0
        root=_node(document_version_id,CanonicalNodeType.DOCUMENT,seq,original_text=filename,normalized_text=filename,parser_confidence=1.0); nodes.append(root); seq+=1
        for i,p in enumerate(doc.paragraphs):
            text=_norm(p.text)
            if not text: continue
            style=(p.style.name or "") if p.style else ""
            if style.lower().startswith("heading"): kind=CanonicalNodeType.HEADING
            elif "list" in style.lower(): kind=CanonicalNodeType.LIST
            else: kind=CanonicalNodeType.PARAGRAPH
            nodes.append(_node(document_version_id,kind,seq,parent_node_id=root.structure_node_id,section_path=style or None,original_text=p.text,normalized_text=text,parser_confidence=1.0,source_locator={"paragraph_index":i},metadata={"paragraph_index":i,"style":style})); seq+=1
        for section_idx,section in enumerate(doc.sections):
            for where,container in (("header",section.header),("footer",section.footer)):
                for i,p in enumerate(container.paragraphs):
                    text=_norm(p.text)
                    if text:
                        nodes.append(_node(document_version_id,CanonicalNodeType.PARAGRAPH,seq,parent_node_id=root.structure_node_id,section_path=where,original_text=p.text,normalized_text=text,parser_confidence=1.0,source_locator={"section_index":section_idx,"container":where,"paragraph_index":i},metadata={"section_index":section_idx,"container":where})); seq+=1
        for ti,table in enumerate(doc.tables):
            t=_node(document_version_id,CanonicalNodeType.TABLE,seq,parent_node_id=root.structure_node_id,parser_confidence=1.0,metadata={"table_id":f"table-{ti}","table_index":ti}); nodes.append(t); seq+=1
            for ri,row in enumerate(table.rows):
                for ci,cell in enumerate(row.cells):
                    text=_norm(cell.text)
                    nodes.append(_node(document_version_id,CanonicalNodeType.TABLE_CELL,seq,parent_node_id=t.structure_node_id,original_text=cell.text,normalized_text=text,parser_confidence=1.0,source_locator={"table_index":ti,"row_index":ri,"column_index":ci},metadata={"table_id":f"table-{ti}","table_index":ti,"row_index":ri,"column_index":ci,"header":ri==0})); seq+=1
        image_index=0
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                nodes.append(_node(document_version_id,CanonicalNodeType.IMAGE,seq,parent_node_id=root.structure_node_id,parser_confidence=1.0,source_locator={"embedded_image_index":image_index},metadata={"embedded_image_ref":rel.target_ref})); seq+=1; image_index+=1
        return ParsedDocument(tuple(nodes),native_parse_used=True,detected_language="und")


class XLSXParserAdapter:
    parser_name="openpyxl"
    parser_version="1.0"
    def supports(self,*,mime_type:str,filename:str)->bool:
        return mime_type=="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or filename.lower().endswith(".xlsx")
    def parse(self,*,document_version_id:UUID,content:bytes,filename:str,mime_type:str)->ParsedDocument:
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(content),data_only=False)
        nodes=[]; seq=0
        root=_node(document_version_id,CanonicalNodeType.DOCUMENT,seq,original_text=filename,normalized_text=filename,parser_confidence=1.0); nodes.append(root); seq+=1
        for ws in wb.worksheets:
            sheet=_node(document_version_id,CanonicalNodeType.SHEET,seq,parent_node_id=root.structure_node_id,sheet_name=ws.title,parser_confidence=1.0,metadata={"max_row":ws.max_row,"max_column":ws.max_column}); nodes.append(sheet); seq+=1
            first_nonempty=None
            for ri,row in enumerate(ws.iter_rows(),start=1):
                values=[cell.value for cell in row]
                if first_nonempty is None and any(v not in (None,"") for v in values): first_nonempty=ri
            table=_node(document_version_id,CanonicalNodeType.TABLE,seq,parent_node_id=sheet.structure_node_id,sheet_name=ws.title,parser_confidence=1.0,metadata={"table_id":f"{ws.title}:used-range"}); nodes.append(table); seq+=1
            header_values=[]
            if first_nonempty:
                header_values=[_norm(c.value).casefold() for c in ws[first_nonempty]]
            merged={}
            for rng in ws.merged_cells.ranges:
                for row in ws[rng.coord]:
                    for cell in row: merged[cell.coordinate]=str(rng)
            for ri,row in enumerate(ws.iter_rows(),start=1):
                row_norm=[_norm(c.value).casefold() for c in row]
                repeated=bool(header_values and ri!=first_nonempty and row_norm[:len(header_values)]==header_values)
                for ci,cell in enumerate(row,start=1):
                    value=cell.value
                    text="" if value is None else str(value)
                    formula=isinstance(value,str) and value.startswith("=")
                    meta={"table_id":f"{ws.title}:used-range","row_index":ri,"column_index":ci,"header":ri==first_nonempty or repeated,"repeated_header":repeated,"formula":formula,"merged_range":merged.get(cell.coordinate),"empty":value in (None,"")}
                    if ci==1 and _norm(value).casefold() in {"unit","units"}: meta["unit_row"]=True
                    if ci==1 and _norm(value).casefold() in {"note","notes"}: meta["note_row"]=True
                    nodes.append(_node(document_version_id,CanonicalNodeType.TABLE_CELL,seq,parent_node_id=table.structure_node_id,sheet_name=ws.title,original_text=text,normalized_text=_norm(text),parser_confidence=1.0,source_locator={"sheet_name":ws.title,"row_index":ri,"column_index":ci,"cell":cell.coordinate},metadata=meta)); seq+=1
        return ParsedDocument(tuple(nodes),native_parse_used=True,detected_language="und")


class CSVParserAdapter:
    parser_name="python-csv"
    parser_version="1.0"
    def supports(self,*,mime_type:str,filename:str)->bool:
        return mime_type=="text/csv" or filename.lower().endswith(".csv")
    def parse(self,*,document_version_id:UUID,content:bytes,filename:str,mime_type:str)->ParsedDocument:
        text=content.decode("utf-8-sig"); rows=list(csv.reader(io.StringIO(text))); nodes=[]; seq=0
        root=_node(document_version_id,CanonicalNodeType.DOCUMENT,seq,original_text=filename,normalized_text=filename,parser_confidence=1.0); nodes.append(root); seq+=1
        sheet=_node(document_version_id,CanonicalNodeType.SHEET,seq,parent_node_id=root.structure_node_id,sheet_name="CSV",parser_confidence=1.0); nodes.append(sheet); seq+=1
        table=_node(document_version_id,CanonicalNodeType.TABLE,seq,parent_node_id=sheet.structure_node_id,sheet_name="CSV",parser_confidence=1.0,metadata={"table_id":"csv:0"}); nodes.append(table); seq+=1
        for ri,row in enumerate(rows,start=1):
            for ci,value in enumerate(row,start=1):
                nodes.append(_node(document_version_id,CanonicalNodeType.TABLE_CELL,seq,parent_node_id=table.structure_node_id,sheet_name="CSV",original_text=value,normalized_text=_norm(value),parser_confidence=1.0,source_locator={"sheet_name":"CSV","row_index":ri,"column_index":ci},metadata={"table_id":"csv:0","row_index":ri,"column_index":ci,"header":ri==1})); seq+=1
        return ParsedDocument(tuple(nodes),native_parse_used=True,detected_language="und")


class PPTXParserAdapter:
    parser_name="python-pptx"
    parser_version="1.0"
    def supports(self,*,mime_type:str,filename:str)->bool:
        return mime_type=="application/vnd.openxmlformats-officedocument.presentationml.presentation" or filename.lower().endswith(".pptx")
    def parse(self,*,document_version_id:UUID,content:bytes,filename:str,mime_type:str)->ParsedDocument:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs=Presentation(io.BytesIO(content)); nodes=[]; seq=0
        root=_node(document_version_id,CanonicalNodeType.DOCUMENT,seq,original_text=filename,normalized_text=filename,parser_confidence=1.0); nodes.append(root); seq+=1
        for si,slide in enumerate(prs.slides,start=1):
            sn=_node(document_version_id,CanonicalNodeType.SLIDE,seq,parent_node_id=root.structure_node_id,slide_no=si,parser_confidence=1.0); nodes.append(sn); seq+=1
            shape_candidates=[]; connector_count=0
            for sh_idx,shape in enumerate(slide.shapes):
                bbox=(float(shape.left),float(shape.top),float(shape.width),float(shape.height))
                if shape.shape_type==MSO_SHAPE_TYPE.TABLE:
                    t=_node(document_version_id,CanonicalNodeType.TABLE,seq,parent_node_id=sn.structure_node_id,slide_no=si,bbox=bbox,parser_confidence=1.0,metadata={"table_id":f"slide-{si}-table-{sh_idx}"}); nodes.append(t); seq+=1
                    for ri,row in enumerate(shape.table.rows):
                        for ci,cell in enumerate(row.cells):
                            txt=_norm(cell.text)
                            nodes.append(_node(document_version_id,CanonicalNodeType.TABLE_CELL,seq,parent_node_id=t.structure_node_id,slide_no=si,original_text=cell.text,normalized_text=txt,parser_confidence=1.0,source_locator={"slide_no":si,"table_index":sh_idx,"row_index":ri,"column_index":ci},metadata={"table_id":f"slide-{si}-table-{sh_idx}","row_index":ri,"column_index":ci,"header":ri==0})); seq+=1
                elif shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                    nodes.append(_node(document_version_id,CanonicalNodeType.IMAGE,seq,parent_node_id=sn.structure_node_id,slide_no=si,bbox=bbox,parser_confidence=1.0,source_locator={"slide_no":si,"shape_index":sh_idx})); seq+=1
                else:
                    txt=_norm(getattr(shape,"text",""))
                    if txt:
                        kind=CanonicalNodeType.HEADING if shape==slide.shapes.title else CanonicalNodeType.PARAGRAPH
                        nodes.append(_node(document_version_id,kind,seq,parent_node_id=sn.structure_node_id,slide_no=si,bbox=bbox,original_text=txt,normalized_text=txt,parser_confidence=1.0,source_locator={"slide_no":si,"shape_index":sh_idx},metadata={"shape_type":str(shape.shape_type)})); seq+=1
                        shape_candidates.append((uuid4(),txt,bbox))
                    if str(shape.shape_type).endswith("LINE") or "CONNECTOR" in str(shape.shape_type).upper(): connector_count+=1
            if len(shape_candidates)>=2 and connector_count:
                a,b=shape_candidates[0],shape_candidates[1]
                meta={"candidate_nodes":[{"candidate_node_id":str(a[0]),"label":a[1],"node_type_candidate":"SYSTEM","confidence":0.7},{"candidate_node_id":str(b[0]),"label":b[1],"node_type_candidate":"SYSTEM","confidence":0.7}],"candidate_edges":[{"candidate_edge_id":str(uuid4()),"source_candidate_node_id":str(a[0]),"target_candidate_node_id":str(b[0]),"direction":"UNKNOWN","relation":"CONNECTOR","confidence":0.5}]}
                nodes.append(_node(document_version_id,CanonicalNodeType.DIAGRAM,seq,parent_node_id=sn.structure_node_id,slide_no=si,parser_confidence=0.6,source_locator={"slide_no":si},metadata=meta)); seq+=1
            if getattr(slide,"has_notes_slide",False):
                notes=_norm(slide.notes_slide.notes_text_frame.text)
                if notes:
                    nodes.append(_node(document_version_id,CanonicalNodeType.PARAGRAPH,seq,parent_node_id=sn.structure_node_id,slide_no=si,section_path="speaker-notes",original_text=notes,normalized_text=notes,parser_confidence=1.0,source_locator={"slide_no":si,"speaker_notes":True})); seq+=1
        return ParsedDocument(tuple(nodes),native_parse_used=True,vision_used=any(n.node_type==CanonicalNodeType.DIAGRAM for n in nodes),detected_language="und")


class TextParserAdapter:
    parser_name="plain-text"
    parser_version="1.0"
    def supports(self,*,mime_type:str,filename:str)->bool: return mime_type=="text/plain" or filename.lower().endswith(".txt")
    def parse(self,*,document_version_id:UUID,content:bytes,filename:str,mime_type:str)->ParsedDocument:
        text=content.decode("utf-8"); nodes=[]; root=_node(document_version_id,CanonicalNodeType.DOCUMENT,0,original_text=filename,normalized_text=filename,parser_confidence=1.0); nodes.append(root)
        for i,p in enumerate([p for p in text.splitlines() if p.strip()],start=1):
            nodes.append(_node(document_version_id,CanonicalNodeType.PARAGRAPH,i,parent_node_id=root.structure_node_id,original_text=p,normalized_text=_norm(p),parser_confidence=1.0,source_locator={"line":i}))
        return ParsedDocument(tuple(nodes),native_parse_used=True,detected_language="und")


class ImageParserAdapter:
    parser_name="pillow+vision-adapter"
    parser_version="1.0"
    def __init__(self,vision:VisionDiagramPort|None=None): self.vision=vision
    def supports(self,*,mime_type:str,filename:str)->bool: return mime_type in {"image/png","image/jpeg"} or filename.lower().endswith((".png",".jpg",".jpeg"))
    def parse(self,*,document_version_id:UUID,content:bytes,filename:str,mime_type:str)->ParsedDocument:
        from PIL import Image
        img=Image.open(io.BytesIO(content)); nodes=[]; root=_node(document_version_id,CanonicalNodeType.DOCUMENT,0,original_text=filename,normalized_text=filename,parser_confidence=1.0); nodes.append(root)
        image=_node(document_version_id,CanonicalNodeType.IMAGE,1,parent_node_id=root.structure_node_id,bbox=(0.0,0.0,float(img.width),float(img.height)),parser_confidence=1.0,metadata={"width":img.width,"height":img.height}); nodes.append(image)
        vision_used=False
        if self.vision:
            result=self.vision.analyze(content=content,mime_type=mime_type); vision_used=True
            meta={"candidate_nodes":[{"candidate_node_id":str(n.candidate_node_id),"label":n.label,"node_type_candidate":n.node_type_candidate,"bbox":n.bbox,"confidence":n.confidence} for n in result.nodes],"candidate_edges":[{"candidate_edge_id":str(e.candidate_edge_id),"source_candidate_node_id":str(e.source_candidate_node_id),"target_candidate_node_id":str(e.target_candidate_node_id),"relation":e.relation,"direction":e.direction,"bbox":e.bbox,"confidence":e.confidence} for e in result.edges],"labels":list(result.labels)}
            nodes.append(_node(document_version_id,CanonicalNodeType.DIAGRAM,2,parent_node_id=root.structure_node_id,bbox=image.bbox,parser_confidence=result.confidence,metadata=meta))
        return ParsedDocument(tuple(nodes),native_parse_used=True,vision_used=vision_used,detected_language=None)


def default_native_parsers(*,ocr:OCRPort|None=None,vision:VisionDiagramPort|None=None):
    return [PDFParserAdapter(ocr),DOCXParserAdapter(),XLSXParserAdapter(),CSVParserAdapter(),PPTXParserAdapter(),TextParserAdapter(),ImageParserAdapter(vision)]
